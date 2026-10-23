"""Word, Excel and PowerPoint files (Office Open XML: ``.docx``, ``.xlsx`` and ``.xlsm``, ``.pptx``), read into what
they hold (``office_content``), which the panel beside a chat shows (``dotagent.previews``) and the model reads as
text (``office_text``). A document holds its paragraphs and tables in order; its text marks headings and list items
as Markdown marks them. A workbook holds its sheets and each sheet the cells that hold anything, a formula with the
value Excel last worked out for it. A presentation holds its slides: the text of their shapes, their tables, their
charts' series and their notes.

The formats of before 2007 (``.doc``, ``.xls``, ``.ppt``) and files locked with a password are not read: they are not
the zip packages of Office Open XML, and the error says how to save one that is.
"""

import datetime as dt
import io
import zipfile
import zlib
from collections.abc import Iterator, Sequence
from contextlib import closing
from itertools import groupby
from operator import attrgetter
from pathlib import PurePath
from typing import Annotated, Literal

import docx
import openpyxl
import pptx
from docx.document import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from pptx.chart.chart import Chart
from pptx.chart.plot import _BasePlot
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError
from pptx.presentation import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.slide import Slide
from pptx.text.text import TextFrame
from pydantic import Field, ValidationError

from dotagent.core.errors import InvalidValueError
from dotagent.core.types import FrozenModel

type OfficeKind = Literal["document", "workbook", "presentation"]

OFFICE_SUFFIXES: dict[str, OfficeKind] = {
    ".docx": "document",
    ".xlsx": "workbook",
    ".xlsm": "workbook",
    ".pptx": "presentation",
}
OFFICE_MEDIA_TYPES: dict[OfficeKind, str] = {
    "document": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "workbook": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "presentation": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}
_NAMES: dict[OfficeKind, str] = {
    "document": "Word document",
    "workbook": "Excel workbook",
    "presentation": "presentation",
}
# The formats of before 2007, and the one to save each as.
_OLD_SUFFIXES: dict[str, str] = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}
_ZIP_MAGIC: bytes = b"PK\x03\x04"
# The first bytes of an OLE compound file: an Office file of before 2007, or one of Office Open XML locked with a
# password.
_OLE_MAGIC: bytes = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
# The cells of a workbook the model reads at most; a sheet larger than that is read in part, and the text says where
# it stops.
MAX_CELLS: int = 100_000
# The places of the rows walked at most for each cell that may be read (item 556): a row is read as wide as its last
# cell, which may hold nothing, so a small file whose rows each end in an empty cell of the last column (XFD) kept a
# worker busy for minutes. A sheet of real data reaches the cells it may read long before.
_PLACES_PER_CELL: int = 50
# The ways a part of an Office file is packed that every program writing one uses: one packed another way, or locked,
# is refused as the file is (item 556), rather than left to fail inside the reader of its kind.
_PACKED: frozenset[int] = frozenset({zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED})
_LOCKED_PART: int = 0x1


class OfficeTable(FrozenModel):
    """A table of a document or of a slide: the text of its cells row by row, the first row its head."""

    kind: Literal["table"]
    rows: tuple[tuple[str, ...], ...]


class DocumentParagraph(FrozenModel):
    """A paragraph of a document that holds text: a heading of level ``heading`` (1 for the title), else an item of a
    list when ``listed``, else plain text."""

    kind: Literal["paragraph"]
    text: str
    heading: int | None
    listed: bool


type DocumentBlock = Annotated[DocumentParagraph | OfficeTable, Field(discriminator="kind")]


class DocumentContent(FrozenModel):
    """A Word document: its paragraphs that hold text and its tables, in order."""

    kind: Literal["document"]
    blocks: tuple[DocumentBlock, ...]


class SheetCell(FrozenModel):
    """A cell that holds anything, by its row and its column (from 1): the value it shows, as text, and its formula
    (``None`` for a value). A formula's value is ``None`` while no program that calculates has worked it out and
    saved it."""

    row: int
    column: int
    value: str | None
    formula: str | None


class Sheet(FrozenModel):
    """A sheet: its title, whether it is shown, and its cells row by row. ``cut_at_row`` is the row reading stopped at
    when the file holds more cells than are read, else ``None``."""

    title: str
    state: Literal["visible", "hidden", "veryHidden"]
    cells: tuple[SheetCell, ...]
    cut_at_row: int | None


class WorkbookContent(FrozenModel):
    """A workbook of ``sheet_count`` sheets, read in order to at most ``max_cells`` cells: ``sheets`` are those read,
    the last of them cut where the cells ran out when they did."""

    kind: Literal["workbook"]
    sheet_count: int
    max_cells: int
    sheets: tuple[Sheet, ...]


class SlideLine(FrozenModel):
    """A paragraph of a shape's text and its level of indent (0 for none)."""

    level: int
    text: str


class SlideText(FrozenModel):
    """The paragraphs of a shape that hold text."""

    kind: Literal["text"]
    lines: tuple[SlideLine, ...]


class ChartPoint(FrozenModel):
    """A value of a series, as text (``None`` when it is blank), and its category (``None`` past the categories)."""

    category: str | None
    value: str | None


class ChartSeries(FrozenModel):
    name: str
    points: tuple[ChartPoint, ...]


class SlideChart(FrozenModel):
    """A chart: its title (empty when it has none), the kinds of its plots and its series. ``kinds`` is ``None`` for a
    chart of a kind python-pptx does not read (3D bars or lines, stock or surface), whose series are not read."""

    kind: Literal["chart"]
    title: str
    kinds: tuple[str, ...] | None
    series: tuple[ChartSeries, ...]


type SlidePart = Annotated[SlideText | OfficeTable | SlideChart, Field(discriminator="kind")]


class SlideContent(FrozenModel):
    """A slide: its title (empty when it has none), what its other shapes hold in order, and its notes."""

    title: str
    parts: tuple[SlidePart, ...]
    notes: str


class PresentationContent(FrozenModel):
    kind: Literal["presentation"]
    slides: tuple[SlideContent, ...]


type OfficeContent = Annotated[DocumentContent | WorkbookContent | PresentationContent, Field(discriminator="kind")]


def office_kind(name: str, head: bytes) -> OfficeKind | None:
    """What kind of Office file ``name`` is, by its suffix, when its first bytes are those of a zip package; ``None``
    for any other file."""
    kind: OfficeKind | None = OFFICE_SUFFIXES.get(PurePath(name).suffix.lower())
    return kind if kind is not None and head.startswith(_ZIP_MAGIC) else None


def unreadable_office_file(name: str, head: bytes) -> str | None:
    """Why ``name``, an Office file by its suffix, cannot be read, or ``None`` when it is no such file: one of before
    2007, or one locked with a password."""
    suffix: str = PurePath(name).suffix.lower()
    newer: str | None = _OLD_SUFFIXES.get(suffix)
    if newer is not None:
        return f"a {suffix} file is of the format of before 2007; save it as {newer}"
    if suffix in OFFICE_SUFFIXES and head.startswith(_OLE_MAGIC):
        return f"the {suffix} file is locked with a password; save a copy without one"
    return None


def names_office_file(name: str) -> bool:
    """Whether ``name`` has the suffix of a Word, Excel or PowerPoint file, of Office Open XML or of before 2007."""
    suffix: str = PurePath(name).suffix.lower()
    return suffix in OFFICE_SUFFIXES or suffix in _OLD_SUFFIXES


def _package_problem(data: bytes) -> str | None:
    """Why the parts of the zip package ``data`` cannot be read, or ``None`` when each is packed as Office packs its
    parts and none is locked."""
    with zipfile.ZipFile(io.BytesIO(data)) as package:
        for part in package.infolist():
            if part.flag_bits & _LOCKED_PART:
                return f"its part {part.filename} is locked"
            if part.compress_type not in _PACKED:
                return f"its part {part.filename} is packed in a way Office does not use"
    return None


def office_content(kind: OfficeKind, data: bytes, max_cells: int) -> OfficeContent:
    """What an Office file of ``kind`` whose bytes are ``data`` holds, a workbook read to at most ``max_cells`` cells.
    ``InvalidValueError`` when it cannot be read: a package that is no zip, lacks a part or has one packed in another
    way or locked, a part whose packed data is damaged (``zlib.error``, or ``EOFError`` when it is cut short), or a part
    whose XML is broken (the errors of lxml and of ElementTree are ``SyntaxError``s)."""
    try:
        problem: str | None = _package_problem(data)
        if problem is not None:
            raise InvalidValueError(f"the {_NAMES[kind]} cannot be read: {problem}", {"kind": kind})
        match kind:
            case "document":
                return read_document(docx.Document(io.BytesIO(data)))
            case "workbook":
                with (
                    closing(openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)) as formulas,
                    closing(openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)) as values,
                ):
                    return read_workbook(_whole(formulas), _whole(values), max_cells)
            case "presentation":
                return read_presentation(pptx.Presentation(io.BytesIO(data)))
    except ValidationError:
        # A model this module builds that does not validate is its own fault, not the file's.
        raise
    except (
        zipfile.BadZipFile,
        zlib.error,
        EOFError,
        KeyError,
        ValueError,
        SyntaxError,
        DocxPackageNotFoundError,
        PptxPackageNotFoundError,
    ) as error:
        raise InvalidValueError(
            f"the {_NAMES[kind]} cannot be read: {error}", {"kind": kind, "error": type(error).__name__}
        ) from error


def office_text(kind: OfficeKind, data: bytes) -> str:
    """The text of an Office file of ``kind`` whose bytes are ``data``, for the model; ``InvalidValueError`` as
    :func:`office_content` raises it."""
    content: OfficeContent = office_content(kind, data, MAX_CELLS)
    match content:
        case DocumentContent():
            return document_text(content)
        case WorkbookContent():
            return workbook_text(content)
        case PresentationContent():
            return presentation_text(content)


def _has_cells(rows: Sequence[Sequence[str]]) -> bool:
    return any(len(row) > 0 for row in rows)


def _cell_text(text: str) -> str:
    """A table cell's text on one line, with the pipes that would end its cell escaped."""
    return " ".join(text.split()).replace("|", "\\|")


def table_lines(rows: Sequence[Sequence[str]]) -> list[str]:
    """A table as Markdown: its first row as the head, every row as wide as the widest."""
    width: int = max((len(row) for row in rows), default=0)
    if width == 0:
        return []
    filled: list[list[str]] = [[_cell_text(cell) for cell in row] + [""] * (width - len(row)) for row in rows]
    head, *body = filled
    return [
        f"| {' | '.join(head)} |",
        f"|{'|'.join(['---'] * width)}|",
        *(f"| {' | '.join(row)} |" for row in body),
    ]


# ----- documents ---------------------------------------------------------------------------------


def _heading_level(style: str) -> int | None:
    """The level of a heading by its style's name: ``Title`` is 1, ``Heading 2`` is 2; ``None`` for no heading."""
    if style == "Title":
        return 1
    if style.startswith("Heading "):
        number: str = style.removeprefix("Heading ")
        return int(number) if number.isdigit() else None
    return None


def _paragraph(paragraph: Paragraph) -> DocumentParagraph | None:
    text: str = paragraph.text.strip()
    if text == "":
        return None
    style: str = paragraph.style.name if paragraph.style is not None and paragraph.style.name is not None else ""
    level: int | None = _heading_level(style)
    return DocumentParagraph(kind="paragraph", text=text, heading=level, listed=level is None and "List" in style)


def _word_rows(table: WordTable) -> tuple[tuple[str, ...], ...]:
    return tuple(tuple(cell.text for cell in row.cells) for row in table.rows)


def read_document(document: WordDocument) -> DocumentContent:
    """A Word document's paragraphs that hold text and its tables that hold cells, in order."""
    blocks: list[DocumentBlock] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph: DocumentParagraph | None = _paragraph(block)
            if paragraph is not None:
                blocks.append(paragraph)
        else:
            rows: tuple[tuple[str, ...], ...] = _word_rows(block)
            if _has_cells(rows):
                blocks.append(OfficeTable(kind="table", rows=rows))
    return DocumentContent(kind="document", blocks=tuple(blocks))


def _block_text(block: DocumentBlock) -> str:
    if isinstance(block, OfficeTable):
        return "\n".join(table_lines(block.rows))
    if block.heading is not None:
        return f"{'#' * min(block.heading, 6)} {block.text}"
    return f"- {block.text}" if block.listed else block.text


def document_text(content: DocumentContent) -> str:
    """A Word document's paragraphs and tables in order, a blank line between them."""
    return "\n\n".join(_block_text(block) for block in content.blocks)


# ----- workbooks ---------------------------------------------------------------------------------


def cell_value(value: object) -> str:
    """A value of a cell as the text shows it: a whole number without a fraction, a date without a time of 00:00."""
    match value:
        case bool():
            return "TRUE" if value else "FALSE"
        case float() if value.is_integer():
            return str(int(value))
        case dt.datetime() if value.time() == dt.time():
            return value.date().isoformat()
        case dt.datetime() | dt.date() | dt.time():
            return value.isoformat()
        case str():
            return " ↵ ".join(value.splitlines())
        case _:
            return str(value)


def column_letters(number: int) -> str:
    """The letters of column ``number`` (from 1): ``A``, ..., ``Z``, ``AA``."""
    letters: str = ""
    while number > 0:
        number, rest = divmod(number - 1, 26)
        letters = chr(ord("A") + rest) + letters
    return letters


def _formula(value: object) -> str | None:
    """The formula a cell holds, or ``None`` for a value."""
    if isinstance(value, ArrayFormula):
        return value.text
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _whole(workbook: Workbook) -> Workbook:
    """``workbook``, read only, with every sheet read to its last cell: a sheet read only stops where the size it
    records of itself ends, and a file written by another program can record one too small."""
    for sheet in workbook.worksheets:
        sheet.reset_dimensions()
    return workbook


def _sheet_cell(row: int, column: int, written: object, worked_out: object) -> SheetCell:
    formula: str | None = _formula(written)
    if formula is None:
        return SheetCell(row=row, column=column, value=cell_value(written), formula=None)
    shown: str | None = None if worked_out is None else cell_value(worked_out)
    return SheetCell(row=row, column=column, value=shown, formula=formula)


def read_workbook(formulas: Workbook, values: Workbook, max_cells: int) -> WorkbookContent:
    """The sheets of a workbook in order and their cells, at most ``max_cells`` of them, and at most
    ``_PLACES_PER_CELL`` times as many places of their rows: ``formulas`` is the workbook read for what its cells hold,
    ``values`` for what Excel last worked out."""
    sheets: list[Sheet] = []
    left: int = max_cells
    places: int = max_cells * _PLACES_PER_CELL
    for written, worked_out in zip(formulas.worksheets, values.worksheets, strict=True):
        cells: list[SheetCell] = []
        cut_at_row: int | None = None
        rows: Iterator[tuple[tuple[object, ...], tuple[object, ...]]] = zip(
            written.iter_rows(values_only=True), worked_out.iter_rows(values_only=True), strict=False
        )
        for row_number, (held, shown) in enumerate(rows, 1):
            places -= max(len(held), 1)
            if places < 0:
                cut_at_row = row_number
                break
            row: list[SheetCell] = [
                _sheet_cell(row_number, column, value, shown[column - 1] if column <= len(shown) else None)
                for column, value in enumerate(held, 1)
                if value is not None
            ]
            if len(row) > left:
                cut_at_row = row_number
                break
            left -= len(row)
            cells.extend(row)
        sheets.append(Sheet(title=written.title, state=written.sheet_state, cells=tuple(cells), cut_at_row=cut_at_row))
        if cut_at_row is not None:
            break
    return WorkbookContent(
        kind="workbook", sheet_count=len(formulas.worksheets), max_cells=max_cells, sheets=tuple(sheets)
    )


def _cell_entry(cell: SheetCell) -> str:
    address: str = f"{column_letters(cell.column)}{cell.row}"
    if cell.formula is None:
        return f"{address}: {cell.value}"
    if cell.value is None:
        return f"{address}: {cell.formula} (not worked out yet)"
    return f"{address}: {cell.formula} → {cell.value}"


def workbook_text(content: WorkbookContent) -> str:
    """Every sheet read: a heading with its name, then a line for each row that holds anything, its cells after their
    addresses, and where reading stopped when it did."""
    sections: list[str] = []
    for number, sheet in enumerate(content.sheets, 1):
        hidden: str = "" if sheet.state == "visible" else f", {sheet.state}"
        lines: list[str] = [f"## Sheet {number} of {content.sheet_count}: {sheet.title}{hidden}"]
        lines.extend(
            " | ".join(_cell_entry(cell) for cell in row) for _, row in groupby(sheet.cells, key=attrgetter("row"))
        )
        if sheet.cut_at_row is not None:
            lines.append(
                f"(the workbook holds more than is read: {content.max_cells} cells at most, fewer when its rows run far"
                f" to the right; the rest from row {sheet.cut_at_row} on is not read)"
            )
        elif len(lines) == 1:
            lines.append("(empty)")
        sections.append("\n".join(lines))
    return "\n\n".join(sections)


# ----- presentations -----------------------------------------------------------------------------

# The kinds of chart python-pptx reads, by the class of their plot. A scatter or bubble chart is read for its y values
# alone.
_PLOT_KINDS: dict[str, str] = {
    "AreaPlot": "area",
    "Area3DPlot": "3D area",
    "BarPlot": "bar",
    "BubblePlot": "bubble, y values only",
    "DoughnutPlot": "doughnut",
    "LinePlot": "line",
    "PiePlot": "pie",
    "RadarPlot": "radar",
    "XyPlot": "scatter, y values only",
}


def _chart(chart: Chart) -> SlideChart:
    """A chart's title, kinds and series with their values by category."""
    title: str = (
        chart.chart_title.text_frame.text.strip() if chart.has_title and chart.chart_title.has_text_frame else ""
    )
    try:
        plots: list[_BasePlot] = list(chart.plots)
    except ValueError:
        # python-pptx reads no chart of another kind than those of _PLOT_KINDS.
        return SlideChart(kind="chart", title=title, kinds=None, series=())
    series: list[ChartSeries] = []
    for plot in plots:
        categories: list[str] = [cell_value(category) for category in plot.categories]
        series.extend(
            ChartSeries(
                name=one.name,
                points=tuple(
                    ChartPoint(
                        category=categories[index] if index < len(categories) else None,
                        value=None if value is None else cell_value(value),
                    )
                    for index, value in enumerate(one.values)
                ),
            )
            for one in plot.series
        )
    kinds: tuple[str, ...] = tuple(dict.fromkeys(_PLOT_KINDS[type(plot).__name__] for plot in plots))
    return SlideChart(kind="chart", title=title, kinds=kinds, series=tuple(series))


def _shape_parts(shape: BaseShape) -> list[SlidePart]:
    """What a shape holds: what the shapes of a group hold, a table, a chart, or the paragraphs of its text that hold
    any, with their levels. A picture or a line holds nothing read."""
    if isinstance(shape, GroupShape):
        return [part for inner in shape.shapes for part in _shape_parts(inner)]
    if isinstance(shape, GraphicFrame) and shape.has_table:
        rows: tuple[tuple[str, ...], ...] = tuple(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
        return [OfficeTable(kind="table", rows=rows)] if _has_cells(rows) else []
    if isinstance(shape, GraphicFrame) and shape.has_chart:
        return [_chart(shape.chart)]
    if isinstance(shape, Shape):
        lines: tuple[SlideLine, ...] = tuple(
            SlideLine(level=paragraph.level, text=paragraph.text.strip())
            for paragraph in shape.text_frame.paragraphs
            if paragraph.text.strip()
        )
        return [SlideText(kind="text", lines=lines)] if lines else []
    return []


def _slide(slide: Slide) -> SlideContent:
    title_shape: Shape | None = slide.shapes.title
    title: str = title_shape.text_frame.text.strip() if title_shape is not None else ""
    parts: list[SlidePart] = [
        part
        for shape in slide.shapes
        if title_shape is None or shape.shape_id != title_shape.shape_id
        for part in _shape_parts(shape)
    ]
    notes: str = ""
    if slide.has_notes_slide:
        frame: TextFrame | None = slide.notes_slide.notes_text_frame
        notes = frame.text.strip() if frame is not None else ""
    return SlideContent(title=title, parts=tuple(parts), notes=notes)


def read_presentation(presentation: Presentation) -> PresentationContent:
    """Every slide of a presentation in order."""
    return PresentationContent(kind="presentation", slides=tuple(_slide(slide) for slide in presentation.slides))


def _point_text(point: ChartPoint) -> str:
    shown: str = "(blank)" if point.value is None else point.value
    return shown if point.category is None else f"{point.category} {shown}"


def _chart_lines(chart: SlideChart) -> list[str]:
    """A chart as a line with its title and kind, then a line for each series with its values by category."""
    named: str = f"Chart “{chart.title}”" if chart.title else "Chart"
    if chart.kinds is None:
        return [f"{named}: of a kind that is not read (3D bars or lines, stock or surface)"]
    lines: list[str] = [f"{named} ({', '.join(chart.kinds)}):" if chart.kinds else f"{named}:"]
    lines.extend(
        f"- {series.name}: {', '.join(_point_text(point) for point in series.points)}" for series in chart.series
    )
    return lines


def _part_lines(part: SlidePart) -> list[str]:
    match part:
        case SlideText():
            return [f"{'  ' * line.level}{line.text}" for line in part.lines]
        case OfficeTable():
            return table_lines(part.rows)
        case SlideChart():
            return _chart_lines(part)


def slide_text(slide: SlideContent, number: int) -> str:
    """A slide under a heading with its number and title: the lines of its other shapes, then its notes."""
    lines: list[str] = [f"## Slide {number}" + (f": {slide.title}" if slide.title else "")]
    for part in slide.parts:
        lines.extend(_part_lines(part))
    if slide.notes:
        lines.append(f"Notes: {slide.notes}")
    return "\n".join(lines)


def presentation_text(content: PresentationContent) -> str:
    """Every slide of a presentation: a heading with its number and title, the text of its shapes, its tables and
    charts, and its notes."""
    return "\n\n".join(slide_text(slide, number) for number, slide in enumerate(content.slides, 1))
