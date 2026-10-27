"""The readers of ``dotagent.office`` that need the libraries of the formats (python-docx, openpyxl, python-pptx),
which weigh about 60 MB with numpy, which openpyxl loads: :func:`dotagent.office.office_content` loads this module on
the first Office file it reads, so that a daemon that reads none never loads them (item 729)."""

import io
from collections.abc import Iterator
from contextlib import closing

import docx
import openpyxl
import pptx
from docx.document import Document as WordDocument
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from docx.table import Table as WordTable
from docx.text.paragraph import Paragraph
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from pptx.chart.chart import Chart
from pptx.chart.plot import _BasePlot
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError
from pptx.presentation import Presentation
from pptx.shapes.autoshape import Shape
from pptx.shapes.base import BaseShape
from pptx.shapes.graphfrm import GraphicFrame
from pptx.shapes.group import GroupShape
from pptx.slide import Slide
from pptx.text.text import TextFrame

from dotagent.office import (
    ChartPoint,
    ChartSeries,
    DocumentBlock,
    DocumentContent,
    DocumentParagraph,
    OfficeContent,
    OfficeKind,
    OfficeTable,
    PresentationContent,
    Sheet,
    SheetCell,
    SlideChart,
    SlideContent,
    SlideLine,
    SlidePart,
    SlideText,
    WorkbookContent,
    cell_value,
    has_cells,
)

# The places of the rows walked at most for each cell that may be read (item 556): a row is read as wide as its last
# cell, which may hold nothing, so a small file whose rows each end in an empty cell of the last column (XFD) kept a
# worker busy for minutes. A sheet of real data reaches the cells it may read long before.
_PLACES_PER_CELL: int = 50
# What the libraries raise for a package that lacks the part they need.
PACKAGE_ERRORS: tuple[type[Exception], ...] = (DocxPackageNotFoundError, PptxPackageNotFoundError)


def read_office(kind: OfficeKind, data: bytes, max_cells: int) -> OfficeContent:
    """What an Office file of ``kind`` whose bytes are ``data`` holds; the errors are those
    :func:`dotagent.office.office_content` turns into its own."""
    match kind:
        case "document":
            return read_document(docx.Document(io.BytesIO(data)))
        case "workbook":
            with (
                closing(openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)) as formulas,
                closing(openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)) as values,
            ):
                return read_workbook(_whole(formulas), _whole(values), max_cells)
        case "presentation":
            return read_presentation(pptx.Presentation(io.BytesIO(data)))


# ----- documents ---------------------------------------------------------------------------------


def _heading_level(style: str) -> int | None:
    """The level of a heading by its style's name: ``Title`` is 1, ``Heading 2`` is 2; ``None`` for no heading."""
    if style == "Title":
        return 1
    if style.startswith("Heading "):
        number: str = style.removeprefix("Heading ")
        return int(number) if number.isdigit() else None
    return None


def _paragraph(paragraph: Paragraph) -> DocumentParagraph | None:
    text: str = paragraph.text.strip()
    if text == "":
        return None
    style: str = paragraph.style.name if paragraph.style is not None and paragraph.style.name is not None else ""
    level: int | None = _heading_level(style)
    return DocumentParagraph(kind="paragraph", text=text, heading=level, listed=level is None and "List" in style)


def _word_rows(table: WordTable) -> tuple[tuple[str, ...], ...]:
    return tuple(tuple(cell.text for cell in row.cells) for row in table.rows)


def read_document(document: WordDocument) -> DocumentContent:
    """A Word document's paragraphs that hold text and its tables that hold cells, in order."""
    blocks: list[DocumentBlock] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            paragraph: DocumentParagraph | None = _paragraph(block)
            if paragraph is not None:
                blocks.append(paragraph)
        else:
            rows: tuple[tuple[str, ...], ...] = _word_rows(block)
            if has_cells(rows):
                blocks.append(OfficeTable(kind="table", rows=rows))
    return DocumentContent(kind="document", blocks=tuple(blocks))


# ----- workbooks ---------------------------------------------------------------------------------


def _formula(value: object) -> str | None:
    """The formula a cell holds, or ``None`` for a value."""
    if isinstance(value, ArrayFormula):
        return value.text
    if isinstance(value, str) and value.startswith("="):
        return value
    return None


def _whole(workbook: Workbook) -> Workbook:
    """``workbook``, read only, with every sheet read to its last cell: a sheet read only stops where the size it
    records of itself ends, and a file written by another program can record one too small."""
    for sheet in workbook.worksheets:
        sheet.reset_dimensions()
    return workbook


def _sheet_cell(row: int, column: int, written: object, worked_out: object) -> SheetCell:
    formula: str | None = _formula(written)
    if formula is None:
        return SheetCell(row=row, column=column, value=cell_value(written), formula=None)
    shown: str | None = None if worked_out is None else cell_value(worked_out)
    return SheetCell(row=row, column=column, value=shown, formula=formula)


def read_workbook(formulas: Workbook, values: Workbook, max_cells: int) -> WorkbookContent:
    """The sheets of a workbook in order and their cells, at most ``max_cells`` of them, and at most
    ``_PLACES_PER_CELL`` times as many places of their rows: ``formulas`` is the workbook read for what its cells hold,
    ``values`` for what Excel last worked out."""
    sheets: list[Sheet] = []
    left: int = max_cells
    places: int = max_cells * _PLACES_PER_CELL
    for written, worked_out in zip(formulas.worksheets, values.worksheets, strict=True):
        cells: list[SheetCell] = []
        cut_at_row: int | None = None
        rows: Iterator[tuple[tuple[object, ...], tuple[object, ...]]] = zip(
            written.iter_rows(values_only=True), worked_out.iter_rows(values_only=True), strict=False
        )
        for row_number, (held, shown) in enumerate(rows, 1):
            places -= max(len(held), 1)
            if places < 0:
                cut_at_row = row_number
                break
            row: list[SheetCell] = [
                _sheet_cell(row_number, column, value, shown[column - 1] if column <= len(shown) else None)
                for column, value in enumerate(held, 1)
                if value is not None
            ]
            if len(row) > left:
                cut_at_row = row_number
                break
            left -= len(row)
            cells.extend(row)
        sheets.append(Sheet(title=written.title, state=written.sheet_state, cells=tuple(cells), cut_at_row=cut_at_row))
        if cut_at_row is not None:
            break
    return WorkbookContent(
        kind="workbook", sheet_count=len(formulas.worksheets), max_cells=max_cells, sheets=tuple(sheets)
    )


# ----- presentations -----------------------------------------------------------------------------

# The kinds of chart python-pptx reads, by the class of their plot. A scatter or bubble chart is read for its y values
# alone.
_PLOT_KINDS: dict[str, str] = {
    "AreaPlot": "area",
    "Area3DPlot": "3D area",
    "BarPlot": "bar",
    "BubblePlot": "bubble, y values only",
    "DoughnutPlot": "doughnut",
    "LinePlot": "line",
    "PiePlot": "pie",
    "RadarPlot": "radar",
    "XyPlot": "scatter, y values only",
}


def _chart(chart: Chart) -> SlideChart:
    """A chart's title, kinds and series with their values by category."""
    title: str = (
        chart.chart_title.text_frame.text.strip() if chart.has_title and chart.chart_title.has_text_frame else ""
    )
    try:
        plots: list[_BasePlot] = list(chart.plots)
    except ValueError:
        # python-pptx reads no chart of another kind than those of _PLOT_KINDS.
        return SlideChart(kind="chart", title=title, kinds=None, series=())
    series: list[ChartSeries] = []
    for plot in plots:
        categories: list[str] = [cell_value(category) for category in plot.categories]
        series.extend(
            ChartSeries(
                name=one.name,
                points=tuple(
                    ChartPoint(
                        category=categories[index] if index < len(categories) else None,
                        value=None if value is None else cell_value(value),
                    )
                    for index, value in enumerate(one.values)
                ),
            )
            for one in plot.series
        )
    kinds: tuple[str, ...] = tuple(dict.fromkeys(_PLOT_KINDS[type(plot).__name__] for plot in plots))
    return SlideChart(kind="chart", title=title, kinds=kinds, series=tuple(series))


def _shape_parts(shape: BaseShape) -> list[SlidePart]:
    """What a shape holds: what the shapes of a group hold, a table, a chart, or the paragraphs of its text that hold
    any, with their levels. A picture or a line holds nothing read."""
    if isinstance(shape, GroupShape):
        return [part for inner in shape.shapes for part in _shape_parts(inner)]
    if isinstance(shape, GraphicFrame) and shape.has_table:
        rows: tuple[tuple[str, ...], ...] = tuple(tuple(cell.text for cell in row.cells) for row in shape.table.rows)
        return [OfficeTable(kind="table", rows=rows)] if has_cells(rows) else []
    if isinstance(shape, GraphicFrame) and shape.has_chart:
        return [_chart(shape.chart)]
    if isinstance(shape, Shape):
        lines: tuple[SlideLine, ...] = tuple(
            SlideLine(level=paragraph.level, text=paragraph.text.strip())
            for paragraph in shape.text_frame.paragraphs
            if paragraph.text.strip()
        )
        return [SlideText(kind="text", lines=lines)] if lines else []
    return []


def _slide(slide: Slide) -> SlideContent:
    title_shape: Shape | None = slide.shapes.title
    title: str = title_shape.text_frame.text.strip() if title_shape is not None else ""
    parts: list[SlidePart] = [
        part
        for shape in slide.shapes
        if title_shape is None or shape.shape_id != title_shape.shape_id
        for part in _shape_parts(shape)
    ]
    notes: str = ""
    if slide.has_notes_slide:
        frame: TextFrame | None = slide.notes_slide.notes_text_frame
        notes = frame.text.strip() if frame is not None else ""
    return SlideContent(title=title, parts=tuple(parts), notes=notes)


def read_presentation(presentation: Presentation) -> PresentationContent:
    """Every slide of a presentation in order."""
    return PresentationContent(kind="presentation", slides=tuple(_slide(slide) for slide in presentation.slides))
