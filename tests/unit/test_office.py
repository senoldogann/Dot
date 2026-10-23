"""Word, Excel and PowerPoint files read as text (item 194): a document's headings, lists and tables in order, a
workbook's cells by their addresses with each formula's value, a presentation's slides with their tables, charts and
notes, and the files that cannot be read and why; read_file pages such a file by its lines as it pages a PDF. The
files are made here with the same libraries, and edited in their zip packages into what those libraries do not
write."""

import io
import zipfile
from collections.abc import Callable
from pathlib import Path

import docx
import openpyxl
import pptx
import pytest
from openpyxl.styles import Font
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from dotagent.core.errors import InvalidValueError, ToolExecutionError
from dotagent.office import (
    WorkbookContent,
    office_content,
    office_kind,
    office_text,
    unreadable_office_file,
    workbook_text,
)
from dotagent.tools.base import ToolResult
from dotagent.tools.files import _read_blocking


def _saved(document: object) -> bytes:
    buffer: io.BytesIO = io.BytesIO()
    document.save(buffer)  # type: ignore[attr-defined]
    return buffer.getvalue()


def test_a_document_reads_as_its_headings_lists_paragraphs_and_tables_in_order() -> None:
    document = docx.Document()
    document.add_heading("Budget 2027", level=1)
    document.add_paragraph("The plan for the year.")
    document.add_paragraph("Rent", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, (("Item", "Amount"), ("Rent | flat", "1200")), strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    document.add_heading("Notes", level=2)
    assert office_text("document", _saved(document)) == (
        "# Budget 2027\n\nThe plan for the year.\n\n- Rent\n\n"
        "| Item | Amount |\n|---|---|\n| Rent \\| flat | 1200 |\n\n## Notes"
    )


def test_a_workbook_reads_as_its_cells_with_each_formula_and_the_value_excel_worked_out() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Budget"
    sheet.append(["Item", "Amount", "Paid"])
    sheet.append(["Rent", 1200, True])
    sheet["B4"] = "=SUM(B2:B3)"
    workbook.create_sheet("Empty").sheet_state = "hidden"
    text: str = office_text("workbook", _saved(workbook))
    # A workbook openpyxl saved has no values worked out: Excel or LibreOffice works them out when it saves.
    assert text == (
        "## Sheet 1 of 2: Budget\nA1: Item | B1: Amount | C1: Paid\nA2: Rent | B2: 1200 | C2: TRUE\n"
        "B4: =SUM(B2:B3) (not worked out yet)\n\n## Sheet 2 of 2: Empty, hidden\n(empty)"
    )


def test_a_presentation_reads_as_its_slides_with_their_tables_charts_and_notes() -> None:
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    assert slide.shapes.title is not None
    slide.shapes.title.text = "Quarter"
    body = slide.placeholders[1].text_frame
    body.text = "Sales grew"
    body.add_paragraph().text = "in every region"
    body.paragraphs[1].level = 1
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2"]
    chart_data.add_series("Sales", (10, 20.5))
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), chart_data
    ).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "Sales by quarter"
    slide.notes_slide.notes_text_frame.text = "Say it slowly."
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    for row, cells in zip(table.rows, (("Region", "Sales"), ("North", "12")), strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    assert office_text("presentation", _saved(presentation)) == (
        "## Slide 1: Quarter\nSales grew\n  in every region\nChart “Sales by quarter” (bar):\n- Sales: Q1 10, Q2 20.5\n"
        "| Region | Sales |\n|---|---|\n| North | 12 |\nNotes: Say it slowly."
    )


def test_what_is_and_is_not_an_office_file_that_can_be_read() -> None:
    package: bytes = b"PK\x03\x04" + bytes(16)
    assert office_kind("Report.DOCX", package) == "document"
    assert office_kind("budget.xlsm", package) == "workbook"
    assert office_kind("budget.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1") is None
    assert office_kind("notes.zip", package) is None
    assert unreadable_office_file("old.xls", b"") == "a .xls file is of the format of before 2007; save it as .xlsx"
    assert unreadable_office_file("locked.docx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(8)) == (
        "the .docx file is locked with a password; save a copy without one"
    )
    assert unreadable_office_file("notes.txt", b"") is None
    broken: io.BytesIO = io.BytesIO()
    with zipfile.ZipFile(broken, "w") as package_file:
        package_file.writestr("hello.txt", "not a document")
    with pytest.raises(InvalidValueError, match="the Word document cannot be read"):
        office_text("document", broken.getvalue())


def _rewritten(package: bytes, part: str, old: bytes, new: bytes) -> bytes:
    """``package`` with ``old`` replaced by ``new`` in its ``part``."""
    written: io.BytesIO = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(package)) as source, zipfile.ZipFile(written, "w") as target:
        for item in source.infolist():
            content: bytes = source.read(item)
            target.writestr(item, content.replace(old, new) if item.filename == part else content)
    return written.getvalue()


def test_a_sheet_that_records_too_small_a_size_of_itself_is_read_to_its_last_cell() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in range(1, 4):
        sheet.append([f"row {row}", row])
    saved: bytes = _saved(workbook)
    assert b'<dimension ref="A1:B3"/>' in zipfile.ZipFile(io.BytesIO(saved)).read("xl/worksheets/sheet1.xml")
    cut: bytes = _rewritten(saved, "xl/worksheets/sheet1.xml", b'<dimension ref="A1:B3"/>', b'<dimension ref="A1"/>')
    assert office_text("workbook", cut) == (
        "## Sheet 1 of 1: Sheet\nA1: row 1 | B1: 1\nA2: row 2 | B2: 2\nA3: row 3 | B3: 3"
    )


def test_a_chart_of_a_kind_that_is_not_read_is_named_and_the_rest_of_the_slide_is_read() -> None:
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    assert slide.shapes.title is not None
    slide.shapes.title.text = "Costs"
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1"]
    chart_data.add_series("Costs", (5,))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), chart_data)
    slide.notes_slide.notes_text_frame.text = "Costs fell."
    # A chart of 3D bars, which python-pptx does not read.
    three_d: bytes = _rewritten(_saved(presentation), "ppt/charts/chart1.xml", b"c:barChart>", b"c:bar3DChart>")
    assert office_text("presentation", three_d) == (
        "## Slide 1: Costs\nChart: of a kind that is not read (3D bars or lines, stock or surface)\nNotes: Costs fell."
    )


def test_read_file_pages_a_workbook_by_its_lines_and_refuses_one_of_before_2007(tmp_path: Path) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in range(1, 31):
        sheet.append([f"Entry {row}", row])
    ledger: Path = tmp_path / "ledger.xlsx"
    ledger.write_bytes(_saved(workbook))
    result: ToolResult = _read_blocking(ledger, 11, 3, 20_000, 1568, 80)
    assert result.text.splitlines() == [
        "    11\tA10: Entry 10 | B10: 10",
        "    12\tA11: Entry 11 | B11: 11",
        "    13\tA12: Entry 12 | B12: 12",
        "[showing lines 11-13; pass offset=14 to continue]",
    ], result.text
    old: Path = tmp_path / "ledger.xls"
    old.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504))
    with pytest.raises(ToolExecutionError) as refused:
        _read_blocking(old, 1, 10, 20_000, 1568, 80)
    assert refused.value.message == "a .xls file is of the format of before 2007; save it as .xlsx"


def _hostile_workbook(rows: int, *, valued: bool) -> bytes:
    """A workbook whose rows each end in a cell of the last column, XFD, which holds a value or only a style: a row is
    read as wide as its last cell, so each costs 16,384 places however little the file weighs (item 556)."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in range(1, rows + 1):
        cell = sheet.cell(row=row, column=16_384, value="x" if valued else None)
        cell.font = Font(bold=True)
    return _saved(workbook)


@pytest.mark.parametrize("valued", [True, False])
def test_a_workbook_whose_rows_run_to_the_last_column_is_read_in_part_whatever_its_cells_hold(valued: bool) -> None:
    hostile: bytes = _hostile_workbook(3, valued=valued)
    # Ten cells may be read, and fifty places for each: the first row alone spans more.
    content = office_content("workbook", hostile, 10)
    assert isinstance(content, WorkbookContent)
    assert [sheet.cut_at_row for sheet in content.sheets] == [1]
    assert workbook_text(content) == (
        "## Sheet 1 of 1: Sheet\n(the workbook holds more than is read: 10 cells at most, fewer when its rows run far "
        "to the right; the rest from row 1 on is not read)"
    )


def test_a_workbook_of_real_data_is_read_whole_however_wide_its_rows_are_for_the_cells_they_hold() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in range(1, 201):
        sheet.append([f"row {row}", row, None, None, row * 2])
    content = office_content("workbook", _saved(workbook), 1_000)
    assert isinstance(content, WorkbookContent)
    assert [(sheet.cut_at_row, len(sheet.cells)) for sheet in content.sheets] == [(None, 600)]


def _repacked(package: bytes, part: str, change: Callable[[zipfile.ZipInfo], None]) -> bytes:
    """``package`` with the entry of ``part`` in its central directory changed by ``change`` once its data is written
    (the packed bytes stay as they are)."""
    written: io.BytesIO = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(package)) as source, zipfile.ZipFile(written, "w") as target:
        for item in source.infolist():
            target.writestr(item, source.read(item), compress_type=item.compress_type)
        change(next(entry for entry in target.filelist if entry.filename == part))
    return written.getvalue()


def _damaged(package: bytes, part: str) -> bytes:
    """``package`` with the first bytes of the packed data of ``part`` set to 0xFF, a block type deflate has none of."""
    with zipfile.ZipFile(io.BytesIO(package)) as source:
        info: zipfile.ZipInfo = source.getinfo(part)
    start: int = info.header_offset + 30 + len(info.filename.encode()) + len(info.extra)
    return package[:start] + b"\xff" * 8 + package[start + 8 :]


def _packed_another_way(entry: zipfile.ZipInfo) -> None:
    entry.compress_type = 9  # Deflate64, which zipfile does not read


def _locked(entry: zipfile.ZipInfo) -> None:
    entry.flag_bits |= 0x1


@pytest.mark.parametrize("damage", ["damaged", "other_packing", "locked"])
def test_a_package_with_a_part_that_cannot_be_unpacked_is_refused_as_a_file_that_cannot_be_read(damage: str) -> None:
    document = docx.Document()
    document.add_paragraph("A letter.")
    package: bytes = _saved(document)
    part: str = "word/document.xml"
    broken: bytes
    match damage:
        case "damaged":
            broken = _damaged(package, part)
        case "other_packing":
            broken = _repacked(package, part, _packed_another_way)
        case _:
            broken = _repacked(package, part, _locked)
    with pytest.raises(InvalidValueError, match="the Word document cannot be read"):
        office_text("document", broken)
