"""Office files as macOS's Quick Look draws them (item 268), with the real ``qlmanage``: a document and a workbook made
here are drawn once per version, a workbook a page per sheet; a drawing's parts are served by their names and nothing
else is; and a file Quick Look cannot draw says so. No network, no model and no browser: the panel's frame is in
``test_file_panel.py``."""

import io
from pathlib import Path

import docx
import httpx
import openpyxl
import pytest
from fastapi import FastAPI

from dotagent.core.errors import NotFoundError, PreviewError
from dotagent.previews import quicklook
from dotagent.previews.quicklook import PAGE, QuickLook, open_quick_look, part_path
from dotagent.server.routes_files import looks_router


def _document(heading: str) -> bytes:
    document = docx.Document()
    document.add_heading(heading, level=1)
    document.add_paragraph("The plan for the year.")
    out: io.BytesIO = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _workbook() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Budget"
    sheet.append(["Item", "Amount"])
    sheet.append(["Rent", 1200])
    workbook.create_sheet("Notes")["A1"] = "Paid on the first"
    out: io.BytesIO = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _written(path: Path, data: bytes) -> Path:
    path.write_bytes(data)
    return path


def _text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _changed_at(path: Path) -> int:
    return path.stat().st_mtime_ns


def _names(folder: Path) -> list[str]:
    return sorted(entry.name for entry in folder.iterdir())


def _sheet_pages(drawing: Path) -> list[str]:
    return [_text(part) for part in drawing.glob("Attachment*.html")]


def _script(drawing: Path) -> str:
    return next(part.name for part in drawing.glob("*.js"))


async def test_a_file_is_drawn_once_a_version_and_a_workbook_a_page_a_sheet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    # Quick Look on a busy runner took longer than the thirty seconds the product gives it (item 540).
    monkeypatch.setattr(quicklook, "DRAW_TIMEOUT_S", 120.0)
    looks: QuickLook = open_quick_look(tmp_path / "looks")
    minutes: Path = _written(tmp_path / "minutes.docx", _document("Budget 2027"))
    drawing: str = await looks.look(minutes)
    page: Path = looks.folder / drawing / PAGE
    assert "Budget 2027" in _text(page)
    drawn_at: int = _changed_at(page)
    assert await looks.look(minutes) == drawing
    assert _changed_at(page) == drawn_at

    # A new version is a new drawing; nothing is left of the work in between.
    _written(minutes, _document("Budget 2028"))
    redrawn: str = await looks.look(minutes)
    assert redrawn != drawing
    assert "Budget 2028" in _text(looks.folder / redrawn / PAGE)
    assert _names(looks.folder) == sorted([drawing, redrawn])

    budget: Path = _written(tmp_path / "budget.xlsx", _workbook())
    sheets: Path = looks.folder / await looks.look(budget)
    tabs: str = _text(sheets / PAGE)
    assert "Budget" in tabs and "Notes" in tabs
    pages: list[str] = _sheet_pages(sheets)
    assert any("Rent" in sheet for sheet in pages) and any("Paid on the first" in sheet for sheet in pages)

    # A daemon that starts again keeps none of the last one's drawings, and its ids are its own.
    again: QuickLook = open_quick_look(tmp_path / "looks")
    assert _names(again.folder) == []
    assert await again.look(minutes) != redrawn


async def test_a_drawings_parts_are_served_by_their_names_and_nothing_else(tmp_path: Path) -> None:
    looks: QuickLook = open_quick_look(tmp_path / "looks")
    drawing: str = await looks.look(_written(tmp_path / "budget.xlsx", _workbook()))
    _written(tmp_path / "secret.txt", b"not a part")
    app: FastAPI = FastAPI()
    app.include_router(looks_router(looks))
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://127.0.0.1") as client:
        served: httpx.Response = await client.get(f"/looks/{drawing}/{PAGE}")
        assert served.status_code == 200
        assert served.headers["content-type"] == "text/html; charset=utf-8"
        assert "Budget" in served.text
        script: httpx.Response = await client.get(f"/looks/{drawing}/{_script(looks.folder / drawing)}")
        assert script.headers["content-type"] == "text/javascript; charset=utf-8"
        # The app answers these with 404 (NotFoundError).
        for address in (f"/looks/{'0' * 32}/{PAGE}", f"/looks/{drawing}/Missing.html", f"/looks/{drawing}/.{PAGE}"):
            with pytest.raises(NotFoundError):
                await client.get(address)
    for drawn, name in ((drawing, "../secret.txt"), (drawing, ".hidden.html"), ("..", "secret.txt"), (drawing, "")):
        with pytest.raises(NotFoundError):
            part_path(looks.folder, drawn, name)


async def test_a_file_quick_look_cannot_draw_says_so_and_leaves_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    looks: QuickLook = open_quick_look(tmp_path / "looks")
    minutes: Path = _written(tmp_path / "minutes.docx", _document("Budget 2027"))
    monkeypatch.setattr(quicklook, "QUICK_LOOK", "/usr/bin/false")
    with pytest.raises(PreviewError, match="Quick Look could not draw this file"):
        await looks.look(minutes)
    assert _names(looks.folder) == []
