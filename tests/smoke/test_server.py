"""The local server end to end, against a real runtime: the guard, the JSON API and the event streams.

Requests reach the application through an ASGI transport, except the tests of live streams, which go through a real
uvicorn server on a free port. The runs use the real local Ollama server and a cloud model; the tests that need it are
skipped when it is not there. The notes are embedded by the local ``bge-m3``.
"""

import asyncio
import base64
import io
import secrets
import sys
import time
from collections.abc import AsyncIterator
from contextlib import aclosing
from dataclasses import dataclass
from pathlib import Path

import docx
import httpx
import openpyxl
import pptx
import pytest
import pytest_asyncio
from PIL import Image
from playwright.async_api import Browser, FilePayload, Page, async_playwright, expect
from structlog.testing import capture_logs

from dotagent.agent.attachments import load_attachments, message_with_attachments
from dotagent.agent.manager import RunManager
from dotagent.cli_daemon import prove_daemon
from dotagent.config import AppConfig, load_config
from dotagent.core.errors import NotFoundError, ServiceError
from dotagent.core.paths import themes_dir
from dotagent.core.types import JsonObject
from dotagent.desktop.perms import check_permissions
from dotagent.desktop.types import Permissions
from dotagent.llm.types import ChatProvider, ImageData, Message
from dotagent.runtime import Runtime, build_runtime, close_runtime
from dotagent.server.app import create_app
from dotagent.server.context import ServerContext
from dotagent.server.daemon import DotServer, bind_listener, serving_config
from dotagent.server.guard import daemon_proof
from dotagent.server.links import ui_url
from dotagent.server.views import PermissionsView
from dotagent.store.endings import RESULT_RESTARTED, RESULT_STOPPED
from dotagent.store.models import ApprovalRecord, EventRecord, RunRecord, ThreadRecord, ThreadSummary

pytestmark = pytest.mark.asyncio(loop_scope="module")

OLLAMA_ROOT: str = "http://localhost:11434"
MODEL: str = "gpt-oss:20b-cloud"
MODEL_REF: str = f"ollama:{MODEL}"
EMBED_MODEL: str = "bge-m3"
TOKEN: str = f"MARKER-{secrets.token_hex(8)}"
PORT: int = 18765
BASE: str = f"http://127.0.0.1:{PORT}"
# The program macOS sees running this process; the daemon of the tests runs in it.
PROGRAM: str = str(Path(sys.executable).resolve())
RUN_TIMEOUT_S: float = 240.0
POLL_S: float = 0.25
SSE_TIMEOUT_S: float = 120.0
# Open connections are cut after three seconds; a server that ends its streams stops well before that.
STOP_WITHIN_S: float = 2.0
# The variable of the Anthropic key in the configuration of the tests: no environment or Keychain has one of that name,
# so the tests of the keys never meet a real key.
KEY_VARIABLE: str = "DOT_SMOKE_ANTHROPIC_KEY"
FAKE_KEY: str = f"sk-ant-smoke-{secrets.token_hex(12)}"
PNG_1X1: bytes = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4890000000d49444154789c6360f8cfc0f01f0005000201"
    "a5f645400000000049454e44ae426082"
)


def _tags() -> list[str]:
    try:
        answer: httpx.Response = httpx.get(f"{OLLAMA_ROOT}/api/tags", timeout=2.0)
    except httpx.HTTPError:
        return []
    return [model["name"] for model in answer.json()["models"]] if answer.status_code == 200 else []


requires_model = pytest.mark.skipif(MODEL not in _tags(), reason=f"needs the local Ollama server with {MODEL}")
requires_embedder = pytest.mark.skipif(
    not any(name.startswith(EMBED_MODEL) for name in _tags()),
    reason=f"needs the local Ollama server with {EMBED_MODEL}",
)


@dataclass(frozen=True)
class SseMessage:
    id: int | None
    kind: str
    data: EventRecord


def _field_of(line: str) -> tuple[str, str]:
    name, _, value = line.partition(":")
    return name, value.removeprefix(" ")


def _message_of(fields: dict[str, str]) -> SseMessage | None:
    """The message a block of fields makes up; a block of comments only (a keepalive) makes up none."""
    if "event" not in fields:
        return None
    identifier: int | None = int(fields["id"]) if "id" in fields else None
    return SseMessage(identifier, fields["event"], EventRecord.model_validate_json(fields["data"]))


def parse_sse(text: str) -> list[SseMessage]:
    """The messages of a whole event stream."""
    blocks: list[dict[str, str]] = [
        dict(_field_of(line) for line in block.splitlines() if line and not line.startswith(":"))
        for block in text.split("\n\n")
    ]
    return [message for message in map(_message_of, blocks) if message is not None]


async def follow_sse(response: httpx.Response) -> AsyncIterator[SseMessage]:
    """The messages of an event stream as they arrive."""
    fields: dict[str, str] = {}
    async for line in response.aiter_lines():
        if line == "":
            message: SseMessage | None = _message_of(fields)
            fields = {}
            if message is not None:
                yield message
        elif not line.startswith(":"):
            name, value = _field_of(line)
            fields[name] = value


@pytest.fixture(scope="module")
def home(tmp_path_factory: pytest.TempPathFactory) -> Path:
    path: Path = tmp_path_factory.mktemp("home")
    (path / "config.toml").write_text(
        f'[models]\nmain = "{MODEL_REF}"\n\n[policy]\npreset = "cautious"\napproval_timeout_s = 120.0\n\n'
        f'[providers.anthropic]\napi_key_env = "{KEY_VARIABLE}"\n',
        encoding="utf-8",
    )
    # The data directory of "the Chrome of the user": empty, so no Chrome runs there and the real one is never touched.
    (path / "chrome").mkdir()
    return path


@pytest_asyncio.fixture(scope="module", loop_scope="module")
async def runtime(home: Path) -> AsyncIterator[Runtime]:
    patch: pytest.MonkeyPatch = pytest.MonkeyPatch()
    patch.setenv("DOT_HOME", str(home))
    patch.setenv("DOT_CHROME_DIR", str(home / "chrome"))
    built: Runtime = await build_runtime(load_config(home / "config.toml"))
    try:
        yield built
    finally:
        await close_runtime(built)
        patch.undo()


@pytest_asyncio.fixture(loop_scope="module")
async def manager(runtime: Runtime) -> AsyncIterator[RunManager]:
    built: RunManager = RunManager(runtime)
    try:
        yield built
    finally:
        await built.shutdown()


def _context(runtime: Runtime, manager: RunManager, port: int, stopping: asyncio.Event) -> ServerContext:
    return ServerContext(
        runtime=runtime,
        manager=manager,
        scheduler=None,
        token=TOKEN,
        host="127.0.0.1",
        port=port,
        started_at=time.time(),
        stopping=stopping,
    )


@pytest_asyncio.fixture(loop_scope="module")
async def client(runtime: Runtime, manager: RunManager) -> AsyncIterator[httpx.AsyncClient]:
    app = create_app(_context(runtime, manager, PORT, asyncio.Event()))
    async with httpx.AsyncClient(
        transport=httpx.ASGITransport(app=app), base_url=BASE, headers={"Authorization": f"Bearer {TOKEN}"}
    ) as http:
        yield http


async def _start_server(runtime: Runtime, manager: RunManager) -> tuple[DotServer, asyncio.Task[None], str]:
    """A real uvicorn server on a free port: the server, the task that serves and the address it listens on."""
    stopping: asyncio.Event = asyncio.Event()
    listener = bind_listener("127.0.0.1", 0)
    port: int = listener.getsockname()[1]
    server = DotServer(serving_config(create_app(_context(runtime, manager, port, stopping))), stopping)
    serving: asyncio.Task[None] = asyncio.create_task(server.serve(sockets=[listener]))
    while not server.started:
        assert not serving.done(), "the server stopped while starting"
        await asyncio.sleep(0.05)
    return server, serving, f"http://127.0.0.1:{port}"


@pytest_asyncio.fixture(loop_scope="module")
async def live(runtime: Runtime, manager: RunManager) -> AsyncIterator[httpx.AsyncClient]:
    """A client of a real uvicorn server on a free port, for streams that stay open."""
    server, serving, address = await _start_server(runtime, manager)
    try:
        async with httpx.AsyncClient(
            base_url=address, headers={"Authorization": f"Bearer {TOKEN}"}, timeout=SSE_TIMEOUT_S
        ) as http:
            yield http
    finally:
        server.should_exit = True
        await serving


async def _run_to_the_end(client: httpx.AsyncClient, run_id: str) -> RunRecord:
    deadline: float = time.monotonic() + RUN_TIMEOUT_S
    while True:
        answer: httpx.Response = await client.get(f"/api/runs/{run_id}")
        assert answer.status_code == 200, answer.text
        if not answer.json()["active"]:
            return RunRecord.model_validate(answer.json()["run"])
        assert time.monotonic() < deadline, f"run {run_id} did not finish"
        await asyncio.sleep(POLL_S)


async def _start(client: httpx.AsyncClient, goal: str) -> RunRecord:
    answer: httpx.Response = await client.post("/api/runs", json={"goal": goal})
    assert answer.status_code == 201, answer.text
    return RunRecord.model_validate(answer.json())


async def _pending_approval(client: httpx.AsyncClient, run_id: str) -> ApprovalRecord:
    deadline: float = time.monotonic() + RUN_TIMEOUT_S
    while True:
        answer: httpx.Response = await client.get("/api/approvals", params={"status": "pending"})
        assert answer.status_code == 200, answer.text
        for item in answer.json():
            if item["run_id"] == run_id:
                return ApprovalRecord.model_validate(item)
        detail: httpx.Response = await client.get(f"/api/runs/{run_id}")
        assert detail.status_code == 200, detail.text
        assert detail.json()["active"], f"run {run_id} ended without asking: {detail.json()['run']}"
        assert time.monotonic() < deadline, f"run {run_id} never asked for anything"
        await asyncio.sleep(POLL_S)


async def _events(client: httpx.AsyncClient, run_id: str) -> list[EventRecord]:
    answer: httpx.Response = await client.get(f"/api/runs/{run_id}/events")
    assert answer.status_code == 200, answer.text
    return [EventRecord.model_validate(item) for item in answer.json()]


# ------------------------------------------------------------------------------------------ the guard


async def test_only_the_health_check_and_the_proof_are_open_and_the_rest_needs_the_token(
    client: httpx.AsyncClient,
) -> None:
    anonymous: dict[str, str] = {"Authorization": ""}
    health: httpx.Response = await client.get("/api/health", headers=anonymous)
    assert health.status_code == 200
    assert health.json()["ok"] is True
    nonce: str = secrets.token_hex(32)
    proof: httpx.Response = await client.get("/api/proof", params={"nonce": nonce}, headers=anonymous)
    assert proof.status_code == 200
    assert proof.json() == {"proof": daemon_proof(TOKEN, nonce)}
    for not_a_nonce in ("", "0" * 31, "0" * 129, "AB" * 16, "g" * 32):
        refused_nonce: httpx.Response = await client.get("/api/proof", params={"nonce": not_a_nonce}, headers=anonymous)
        assert refused_nonce.status_code == 422, not_a_nonce
    for path in ("/api/status", "/api/runs", "/api/memories", "/api/stream"):
        refused: httpx.Response = await client.get(path, headers=anonymous)
        assert refused.status_code == 401, path
        assert refused.headers["www-authenticate"] == "Bearer"
        wrong: httpx.Response = await client.get(path, headers={"Authorization": "Bearer nope"})
        assert wrong.status_code == 401, path
    assert (await client.get("/api/status")).status_code == 200


async def test_the_daemon_passes_the_check_the_cli_makes_before_it_sends_the_token(
    runtime: Runtime, manager: RunManager
) -> None:
    server, serving, address = await _start_server(runtime, manager)
    port: int = int(address.rpartition(":")[2])
    try:
        async with httpx.AsyncClient(trust_env=False) as plain:
            await prove_daemon(plain, "127.0.0.1", port, TOKEN)
            with pytest.raises(ServiceError):
                await prove_daemon(plain, "127.0.0.1", port, "another token")
    finally:
        server.should_exit = True
        await serving


async def test_other_hosts_and_other_origins_are_refused(client: httpx.AsyncClient) -> None:
    rebinding: httpx.Response = await client.get("/api/health", headers={"Host": "attacker.example:18765"})
    assert rebinding.status_code == 421
    foreign: httpx.Response = await client.get("/api/status", headers={"Origin": "https://attacker.example"})
    assert foreign.status_code == 403
    same: httpx.Response = await client.get("/api/status", headers={"Origin": BASE})
    assert same.status_code == 200


async def test_pages_carry_strict_headers_and_there_is_no_interactive_documentation(client: httpx.AsyncClient) -> None:
    page: httpx.Response = await client.get("/", headers={"Authorization": ""})
    assert page.status_code == 200
    assert "default-src 'none'" in page.headers["content-security-policy"]
    assert "frame-ancestors 'none'" in page.headers["content-security-policy"]
    assert page.headers["x-content-type-options"] == "nosniff"
    assert (await client.get("/api/status")).headers["cache-control"] == "no-store"
    for path in ("/docs", "/redoc", "/openapi.json", "/api/openapi.json"):
        assert (await client.get(path)).status_code == 404, path


async def test_a_malformed_request_is_answered_with_what_is_wrong(client: httpx.AsyncClient) -> None:
    answer: httpx.Response = await client.post("/api/runs", json={"goal": "", "mode": "nonsense"})
    assert answer.status_code == 422
    body: JsonObject = answer.json()
    assert body["message"] == "the request is not valid"
    assert isinstance(body["context"], dict)
    assert len(body["context"]["errors"]) == 2
    unknown_model: httpx.Response = await client.post("/api/runs", json={"goal": "x", "model": "gpt-5"})
    assert unknown_model.status_code == 422
    assert unknown_model.json()["context"]["providers"] == ["ollama", "openai", "opencode_go", "anthropic", "opencode"]
    unknown_thread: httpx.Response = await client.post("/api/runs", json={"goal": "x", "thread_id": "thr_nope"})
    assert unknown_thread.status_code == 404
    assert (await client.get("/api/runs/run_nope")).status_code == 404
    assert (await client.get("/api/runs/run_nope/events")).status_code == 404
    assert (await client.get("/api/runs/run_nope/stream")).status_code == 404
    assert (await client.post("/api/runs/run_nope/cancel")).status_code == 404
    assert (await client.post("/api/approvals/apr_nope/decision", json={"approve": True})).status_code == 404
    assert (await client.get("/api/runs", params={"limit": 0})).status_code == 422


async def test_an_approval_mode_that_does_not_exist_and_goal_mode_for_a_chat_are_refused(
    client: httpx.AsyncClient,
) -> None:
    unknown_mode: httpx.Response = await client.post("/api/runs", json={"goal": "x", "approvals": "reckless"})
    assert unknown_mode.status_code == 422
    assert unknown_mode.json()["message"] == "the request is not valid"
    chat: httpx.Response = await client.post("/api/runs", json={"goal": "x", "mode": "chat", "goal_mode": True})
    assert chat.status_code == 422
    assert "goal mode" in chat.json()["message"], chat.text


# ------------------------------------------------------------------------------------------ what Dot keeps


async def test_rules_are_listed_changed_and_deleted_through_the_api(client: httpx.AsyncClient) -> None:
    created: httpx.Response = await client.post(
        "/api/rules",
        json={
            "behavior": "ask",
            "tool_glob": "shell",
            "match_regex": "rm -rf /tmp/zzz",
            "description": "ask about zzz",
        },
    )
    assert created.status_code == 201, created.text
    rule_id: str = created.json()["id"]
    assert rule_id in [rule["id"] for rule in (await client.get("/api/rules")).json()]
    off: httpx.Response = await client.patch(f"/api/rules/{rule_id}", json={"enabled": False})
    assert off.status_code == 200
    assert off.json()["enabled"] is False
    bad: httpx.Response = await client.post(
        "/api/rules", json={"behavior": "allow", "tool_glob": "*", "match_regex": "(", "description": "broken"}
    )
    assert bad.status_code == 422
    assert bad.json()["context"]["pattern"] == "("
    deleted: httpx.Response = await client.delete(f"/api/rules/{rule_id}")
    assert deleted.status_code == 200
    assert (await client.patch(f"/api/rules/{rule_id}", json={"enabled": True})).status_code == 404
    assert (await client.delete(f"/api/rules/{rule_id}")).status_code == 404


async def test_schedules_are_created_switched_and_deleted_through_the_api(client: httpx.AsyncClient) -> None:
    created: httpx.Response = await client.post(
        "/api/schedules", json={"name": "morning look", "goal": "Say hello.", "cron": "0 9 * * 1-5"}
    )
    assert created.status_code == 201, created.text
    view: JsonObject = created.json()
    schedule_id: str = view["record"]["id"]
    assert view["when"] == "cron 0 9 * * 1-5"
    assert view["next_local"] is not None
    assert view["record"]["enabled"] is True
    off: httpx.Response = await client.post(f"/api/schedules/{schedule_id}/disable")
    assert off.json()["record"]["enabled"] is False
    assert off.json()["next_local"] is None
    on: httpx.Response = await client.post(f"/api/schedules/{schedule_id}/enable")
    assert on.json()["record"]["enabled"] is True
    assert schedule_id in [item["record"]["id"] for item in (await client.get("/api/schedules")).json()]
    change: JsonObject = {
        "name": "evening look",
        "goal": "Say good evening.",
        "cron": "30 18 * * *",
        "at": None,
        "mode": "research",
        "model": "ollama:gpt-oss:20b-cloud",
        "effort": None,
        "approvals": "yolo",
        "goal_mode": True,
        "timezone": "Europe/Istanbul",
        "enabled": True,
    }
    edited: httpx.Response = await client.patch(f"/api/schedules/{schedule_id}", json=change)
    assert edited.status_code == 200, edited.text
    assert (edited.json()["record"]["name"], edited.json()["when"], edited.json()["zone"]) == (
        "evening look",
        "cron 30 18 * * *",
        "Europe/Istanbul",
    )
    assert edited.json()["record"]["options"]["approvals"] == "yolo"
    assert edited.json()["next_local"].endswith("18:30")
    fetched: httpx.Response = await client.get(f"/api/schedules/{schedule_id}")
    assert fetched.json() == edited.json()
    unknown_zone: httpx.Response = await client.patch(
        f"/api/schedules/{schedule_id}", json={**change, "timezone": "Mars/Base"}
    )
    assert unknown_zone.status_code == 422
    assert (await client.patch(f"/api/schedules/{schedule_id}", json={"name": "only a name"})).status_code == 422
    both: httpx.Response = await client.post(
        "/api/schedules", json={"name": "bad", "goal": "x", "cron": "0 9 * * *", "at": "2031-01-01 09:00"}
    )
    assert both.status_code == 422
    assert (await client.delete(f"/api/schedules/{schedule_id}")).status_code == 200
    assert (await client.delete(f"/api/schedules/{schedule_id}")).status_code == 404


@requires_embedder
async def test_memories_are_saved_found_and_forgotten_through_the_api(client: httpx.AsyncClient) -> None:
    marker: str = f"zorblax-{secrets.token_hex(3)}"
    saved: httpx.Response = await client.post(
        "/api/memories", json={"content": f"The user's tortoise is called {marker}.", "tags": ["pets"]}
    )
    assert saved.status_code == 201, saved.text
    memory_id: str = saved.json()["record"]["id"]
    page: httpx.Response = await client.get("/api/memories", params={"limit": 5})
    assert page.json()["total"] >= 1
    assert memory_id in [item["id"] for item in page.json()["items"]]
    found: httpx.Response = await client.get("/api/memories/search", params={"q": marker})
    assert found.status_code == 200
    assert memory_id in [hit["record"]["id"] for hit in found.json()]
    renamed: str = f"The user's tortoise is now called {marker}-ii."
    edited: httpx.Response = await client.patch(
        f"/api/memories/{memory_id}", json={"content": renamed, "tags": ["pet"]}
    )
    assert edited.status_code == 200, edited.text
    assert (edited.json()["content"], edited.json()["tags"], edited.json()["source_run"]) == (renamed, ["pet"], None)
    refound: httpx.Response = await client.get("/api/memories/search", params={"q": f"{marker}-ii"})
    assert [hit["record"]["content"] for hit in refound.json()][:1] == [renamed]
    assert (await client.delete(f"/api/memories/{memory_id}")).status_code == 200
    assert (await client.patch(f"/api/memories/{memory_id}", json={"content": renamed})).status_code == 404
    assert (await client.delete(f"/api/memories/{memory_id}")).status_code == 404
    assert (await client.post("/api/memories", json={"content": "   "})).status_code == 422


def tiny_pdf(text: str) -> bytes:
    """A one-page PDF with ``text`` in its text layer."""
    stream: bytes = f"BT /F1 24 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out: bytearray = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref: int = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def tiny_png() -> bytes:
    buffer: io.BytesIO = io.BytesIO()
    Image.new("RGB", (40, 30), (200, 30, 30)).save(buffer, format="PNG")
    return buffer.getvalue()


def tiny_office(name: str, text: str) -> bytes:
    """A Word document, an Excel workbook or a presentation, by the suffix of ``name``, that holds ``text``."""
    buffer: io.BytesIO = io.BytesIO()
    match Path(name).suffix:
        case ".docx":
            document = docx.Document()
            document.add_paragraph(text)
            document.save(buffer)
        case ".xlsx":
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            assert sheet is not None
            sheet["A1"] = text
            workbook.save(buffer)
        case ".pptx":
            presentation = pptx.Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            assert slide.shapes.title is not None
            slide.shapes.title.text = text
            presentation.save(buffer)
    return buffer.getvalue()


async def _upload(client: httpx.AsyncClient, name: str, media_type: str, data: bytes) -> httpx.Response:
    return await client.post(
        "/api/attachments",
        json={"name": name, "media_type": media_type, "data_b64": base64.b64encode(data).decode("ascii")},
    )


async def test_attached_files_are_kept_and_reach_the_message_the_model_reads(
    client: httpx.AsyncClient, runtime: Runtime
) -> None:
    code: str = f"walrus-{secrets.token_hex(3)}"
    notes: httpx.Response = await _upload(client, "../notes.py", "", f"SECRET = '{code}'\n".encode())
    assert notes.status_code == 201, notes.text
    assert (notes.json()["kind"], notes.json()["name"]) == ("text", "notes.py")
    assert Path(notes.json()["path"]).parent.parent == runtime.workspace / "attachments"
    pdf: httpx.Response = await _upload(client, "report.pdf", "application/pdf", tiny_pdf(f"PDF says {code}-pdf"))
    assert (pdf.status_code, pdf.json()["kind"]) == (201, "pdf"), pdf.text
    photo: httpx.Response = await _upload(client, "photo.png", "", tiny_png())
    assert (photo.status_code, photo.json()["kind"], photo.json()["media_type"]) == (201, "image", "image/png")
    # Word, Excel and PowerPoint files are read for their text (item 194); one of before 2007 or locked with a
    # password is refused, with what to save instead.
    office_names: tuple[str, ...] = ("plan.docx", "budget.xlsx", "deck.pptx")
    office: list[httpx.Response] = [
        await _upload(client, name, "", tiny_office(name, f"{name} says {code}")) for name in office_names
    ]
    assert [sent.status_code for sent in office] == [201] * 3, [sent.text for sent in office]
    assert [sent.json()["kind"] for sent in office] == ["office"] * 3
    compound: bytes = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + bytes(504)
    old: httpx.Response = await _upload(client, "budget.xls", "application/vnd.ms-excel", compound)
    assert (old.status_code, "save it as .xlsx" in old.text) == (422, True), old.text
    locked: httpx.Response = await _upload(client, "plan.docx", "", compound)
    assert (locked.status_code, "locked with a password" in locked.text) == (422, True), locked.text

    # Refused: a file that is no text, PDF, Office file or image, a broken image, bad base64, an unknown or crafted
    # id.
    assert (await _upload(client, "tool.bin", "application/octet-stream", bytes(range(256)) * 4)).status_code == 422
    assert (await _upload(client, "fake.png", "image/png", b"not an image")).status_code == 422
    bad: httpx.Response = await client.post(
        "/api/attachments", json={"name": "x.txt", "media_type": "text/plain", "data_b64": "@@@"}
    )
    assert bad.status_code == 422
    for wrong in ("att_missing_000000", "../config"):
        started: httpx.Response = await client.post("/api/runs", json={"goal": "hi", "attachments": [wrong]})
        assert started.status_code == 404, started.text

    files = await load_attachments(
        runtime.workspace,
        [notes.json()["id"], pdf.json()["id"], photo.json()["id"], *(sent.json()["id"] for sent in office)],
    )
    message: Message = await message_with_attachments(runtime, "thread_test", None, "Use the files.", files, True)
    assert message.content.startswith("Use the files.")
    assert f"SECRET = '{code}'" in message.content
    assert f"PDF says {code}-pdf" in message.content
    assert all(f"{name} says {code}" in message.content for name in office_names)
    assert notes.json()["path"] in message.content
    assert [image.media_type for image in message.images] == ["image/jpeg"]


async def test_screenshots_are_served_by_their_file_name_only(client: httpx.AsyncClient, runtime: Runtime) -> None:
    names: list[str] = await runtime.store.save_images(
        [ImageData(media_type="image/png", data_b64=base64.b64encode(PNG_1X1).decode("ascii"))]
    )
    assert len(names) == 1
    served: httpx.Response = await client.get(f"/api/images/{names[0]}")
    assert served.status_code == 200
    assert served.headers["content-type"] == "image/png"
    assert served.content == PNG_1X1
    for name in (
        "0" * 24 + ".png",
        "../config.toml",
        "..%2Fconfig.toml",
        "abc.png",
        names[0].upper(),
        "x" * 24 + ".png",
    ):
        assert (await client.get(f"/api/images/{name}")).status_code == 404, name


async def test_the_status_the_models_and_the_checks_describe_this_installation(client: httpx.AsyncClient) -> None:
    status: httpx.Response = await client.get("/api/status")
    assert status.status_code == 200
    body: JsonObject = status.json()
    assert body["port"] == PORT
    assert body["preset"] == "cautious"
    assert (body["approvals"], body["goal_mode"]) == ("ask", False)
    assert body["scheduler_running"] is False
    assert body["models"]["main"] == MODEL_REF
    models: httpx.Response = await client.get("/api/models")
    assert models.status_code == 200
    assert [entry["provider"] for entry in models.json()] == [
        "ollama",
        "openai",
        "opencode_go",
        "anthropic",
        "opencode",
    ]
    checks: httpx.Response = await client.get("/api/doctor")
    assert checks.status_code == 200
    names: list[str] = [check["name"] for check in checks.json()]
    assert "browser launch" not in names
    assert "ollama server" in names
    assert "mcp" in names
    assert ("desktop macos" in names) is body["desktop_enabled"]


async def test_the_browser_mode_is_reported_and_only_switched_when_chrome_is_there(
    client: httpx.AsyncClient, home: Path
) -> None:
    state: httpx.Response = await client.get("/api/browser")
    assert state.status_code == 200
    body: JsonObject = state.json()
    assert (body["mode"], body["connected"]) == ("own", False)
    assert (body["chrome_running"], body["remote_debugging"], body["port"]) == (False, False, None)
    assert body["inspect_page"] == "chrome://inspect/#remote-debugging"
    refused: httpx.Response = await client.post("/api/browser/mode", json={"mode": "attach"})
    assert refused.status_code == 422, refused.text
    assert "Chrome is not running" in refused.json()["message"]
    assert (await client.get("/api/browser")).json()["mode"] == "own"
    assert (await client.get("/api/status")).json()["browser_mode"] == "own"
    assert (await client.post("/api/browser/mode", json={"mode": "chrome"})).status_code == 422
    kept: httpx.Response = await client.post("/api/browser/mode", json={"mode": "own"})
    assert kept.status_code == 200, kept.text
    assert kept.json()["mode"] == "own"
    written: AppConfig = load_config(home / "config.toml")
    assert (written.browser.mode, written.policy.preset) == ("own", "cautious")


async def test_whether_chrome_is_raised_is_reported_switched_live_and_kept(
    client: httpx.AsyncClient, home: Path
) -> None:
    assert (await client.get("/api/browser")).json()["raise_window"] is False
    raised: httpx.Response = await client.post("/api/browser/raise", json={"raise_window": True})
    assert raised.status_code == 200, raised.text
    assert raised.json()["raise_window"] is True
    assert (await client.get("/api/browser")).json()["raise_window"] is True
    assert load_config(home / "config.toml").browser.raise_window is True
    lowered: httpx.Response = await client.post("/api/browser/raise", json={"raise_window": False})
    assert lowered.json()["raise_window"] is False
    assert load_config(home / "config.toml").browser.raise_window is False
    assert (await client.post("/api/browser/raise", json={"raise_window": "maybe"})).status_code == 422


def _desktop_is_fully_permitted() -> bool:
    granted: Permissions = check_permissions()
    return granted.accessibility and granted.post_event and granted.screen_recording


@pytest.mark.skipif(
    not _desktop_is_fully_permitted(), reason="asking macOS for a permission it has not granted shows a dialog"
)
async def test_the_daemon_asks_macos_for_its_own_permissions_and_reports_them(client: httpx.AsyncClient) -> None:
    refused: httpx.Response = await client.post("/api/desktop/permissions", headers={"Authorization": ""})
    assert refused.status_code == 401
    answer: httpx.Response = await client.post("/api/desktop/permissions")
    assert answer.status_code == 200, answer.text
    assert PermissionsView.model_validate(answer.json()) == PermissionsView(
        program=PROGRAM,
        accessibility=True,
        input_events=True,
        screen_recording=True,
        panes=[],
    )


# ------------------------------------------------------------------------------------------ API keys


def _key_view(listing: httpx.Response, name: str) -> JsonObject:
    views: list[JsonObject] = [view for view in listing.json()["keys"] if view["name"] == name]
    assert len(views) == 1, [view["name"] for view in listing.json()["keys"]]
    return views[0]


async def test_a_key_is_saved_shown_by_its_end_only_used_at_once_and_removed(
    client: httpx.AsyncClient, runtime: Runtime, home: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    listing: httpx.Response = await client.get("/api/keys")
    assert listing.status_code == 200, listing.text
    assert [view["provider"] for view in listing.json()["keys"]] == [
        "openai",
        "anthropic",
        "opencode_go",
        "ollama",
        "typesafe",
    ]
    assert listing.json()["env_file"] == str(home / ".env")
    assert _key_view(listing, KEY_VARIABLE) == {
        "provider": "anthropic",
        "name": KEY_VARIABLE,
        "source": None,
        "saved": False,
        "tail": "",
    }
    saved: httpx.Response = await client.post(f"/api/keys/{KEY_VARIABLE}", json={"value": FAKE_KEY})
    assert saved.status_code == 200
    assert saved.json() == {
        "provider": "anthropic",
        "name": KEY_VARIABLE,
        "source": "env_file",
        "saved": True,
        "tail": FAKE_KEY[-4:],
    }
    assert FAKE_KEY not in saved.text
    assert FAKE_KEY not in (await client.get("/api/keys")).text
    assert (home / ".env").stat().st_mode & 0o777 == 0o600
    # A changed key reaches the running daemon: the next request makes a new connector with it.
    first: ChatProvider = await runtime.router.provider("anthropic")
    replaced: httpx.Response = await client.post(f"/api/keys/{KEY_VARIABLE}", json={"value": f"{FAKE_KEY}-new"})
    assert replaced.json()["tail"] == "-new"
    assert await runtime.router.provider("anthropic") is not first
    # The environment of the daemon comes before Dot's .env file.
    with monkeypatch.context() as patch:
        patch.setenv(KEY_VARIABLE, f"sk-ant-environment-{secrets.token_hex(12)}")
        shadowed: JsonObject = _key_view(await client.get("/api/keys"), KEY_VARIABLE)
        assert (shadowed["source"], shadowed["saved"]) == ("environment", True)
    # A value that is not a key is refused without being repeated.
    spaced: str = f"sk-ant {secrets.token_hex(12)}"
    refused: httpx.Response = await client.post(f"/api/keys/{KEY_VARIABLE}", json={"value": spaced})
    assert refused.status_code == 422
    assert spaced not in refused.text
    assert (await client.post("/api/keys/NOT_A_KEY", json={"value": FAKE_KEY})).status_code == 404
    assert (await client.post("/api/keys/NOT_A_KEY/check")).status_code == 404
    removed: httpx.Response = await client.delete(f"/api/keys/{KEY_VARIABLE}")
    assert removed.status_code == 200
    assert (removed.json()["source"], removed.json()["saved"]) == (None, False)
    assert KEY_VARIABLE not in (home / ".env").read_text(encoding="utf-8")
    assert (await client.delete(f"/api/keys/{KEY_VARIABLE}")).status_code == 404


@pytest.mark.internet
async def test_checking_a_key_the_provider_refuses_says_so(client: httpx.AsyncClient) -> None:
    assert (await client.post(f"/api/keys/{KEY_VARIABLE}", json={"value": FAKE_KEY})).status_code == 200
    try:
        checked: httpx.Response = await client.post(f"/api/keys/{KEY_VARIABLE}/check")
    finally:
        await client.delete(f"/api/keys/{KEY_VARIABLE}")
    assert checked.status_code == 200, checked.text
    assert (checked.json()["name"], checked.json()["ok"]) == ("anthropic api", False)
    assert "401" in checked.json()["detail"], checked.json()["detail"]


# ------------------------------------------------------------------------------------------ event streams


async def _stored_run(runtime: Runtime, parent_run_id: str | None) -> RunRecord:
    """A run that only exists in the store, in a thread of its own."""
    thread: ThreadRecord = await runtime.store.create_thread("stream test", "test")
    return await runtime.store.create_run(
        thread.id,
        "goal",
        "task",
        MODEL_REF,
        parent_run_id,
        "test",
        "medium",
        "standard",
        None,
        "ask",
        False,
        None,
        None,
        None,
    )


def _queued(parent_run_id: str | None) -> JsonObject:
    return {"goal": "goal", "mode": "task", "model": MODEL_REF, "parent_run_id": parent_run_id}


async def _tell(runtime: Runtime, story: list[tuple[str, str, JsonObject]]) -> list[EventRecord]:
    """Store the events of a story in order: ``(run id, kind, data)``."""
    return [await runtime.bus.emit(run_id, kind, data, True) for run_id, kind, data in story]


async def test_the_stream_of_a_finished_run_replays_it_with_the_runs_it_delegated_to(
    client: httpx.AsyncClient, runtime: Runtime
) -> None:
    parent: RunRecord = await _stored_run(runtime, None)
    child: RunRecord = await _stored_run(runtime, parent.id)
    stranger: RunRecord = await _stored_run(runtime, None)
    told: list[EventRecord] = await _tell(
        runtime,
        [
            (stranger.id, "run_started", {}),
            (parent.id, "run_queued", _queued(None)),
            (parent.id, "run_started", {}),
            (child.id, "run_queued", _queued(parent.id)),
            (stranger.id, "assistant", {"text": "not for this stream"}),
            (child.id, "run_started", {}),
            (child.id, "assistant", {"text": "from the child"}),
            (child.id, "run_finished", {"status": "done"}),
            (parent.id, "assistant", {"text": "from the parent"}),
            (parent.id, "run_finished", {"status": "done"}),
        ],
    )
    mine: list[EventRecord] = [event for event in told if event.run_id != stranger.id]
    answer: httpx.Response = await asyncio.wait_for(client.get(f"/api/runs/{parent.id}/stream"), 10.0)
    assert answer.status_code == 200
    assert answer.headers["content-type"].startswith("text/event-stream")
    assert answer.headers["cache-control"] == "no-store"
    assert [(item.id, item.data.run_id, item.kind) for item in parse_sse(answer.text)] == [
        (event.id, event.run_id, event.kind) for event in mine
    ]
    resumed: httpx.Response = await client.get(f"/api/runs/{parent.id}/stream", params={"after": mine[2].id})
    assert [item.id for item in parse_sse(resumed.text)] == [event.id for event in mine[3:]]
    detail: httpx.Response = await client.get(f"/api/runs/{parent.id}")
    assert [run["id"] for run in detail.json()["children"]] == [child.id]
    assert detail.json()["active"] is False


async def test_the_stream_of_a_run_that_nothing_is_running_ends_after_what_it_stored(
    client: httpx.AsyncClient, runtime: Runtime
) -> None:
    run: RunRecord = await _stored_run(runtime, None)
    told: list[EventRecord] = await _tell(runtime, [(run.id, "run_queued", _queued(None)), (run.id, "run_started", {})])
    answer: httpx.Response = await asyncio.wait_for(client.get(f"/api/runs/{run.id}/stream"), 10.0)
    assert [item.id for item in parse_sse(answer.text)] == [event.id for event in told]


async def test_only_a_run_the_daemon_cut_off_is_offered_to_be_resumed(
    client: httpx.AsyncClient, runtime: Runtime
) -> None:
    cut: RunRecord = await _stored_run(runtime, None)
    await runtime.store.finish_run(cut.id, "failed", RESULT_RESTARTED)
    stopped: RunRecord = await _stored_run(runtime, None)
    await runtime.store.finish_run(stopped.id, "cancelled", RESULT_STOPPED)
    child: RunRecord = await _stored_run(runtime, cut.id)
    await runtime.store.finish_run(child.id, "failed", RESULT_RESTARTED)
    offered: list[bool] = []
    for run in (cut, stopped, child):
        detail: httpx.Response = await client.get(f"/api/runs/{run.id}")
        assert detail.status_code == 200, detail.text
        offered.append(detail.json()["resumable"])
    assert offered == [True, False, False]
    # Resuming is refused with what the run is, and nothing starts.
    refused: httpx.Response = await client.post(f"/api/runs/{stopped.id}/resume")
    assert refused.status_code == 409, refused.text
    assert refused.json()["context"]["result"] == RESULT_STOPPED
    assert [run.id for run in await runtime.store.list_runs(1, stopped.thread_id)] == [stopped.id]
    assert (await client.post("/api/runs/run_missing/resume")).status_code == 404


async def test_the_lifecycle_stream_carries_only_the_moments_a_client_refetches_for(
    live: httpx.AsyncClient, runtime: Runtime
) -> None:
    run: RunRecord = await _stored_run(runtime, None)

    async def chatter() -> None:
        while True:
            await runtime.bus.emit(run.id, "assistant", {"text": "filtered"}, False)
            await runtime.bus.emit(run.id, "run_started", {}, False)
            await asyncio.sleep(0.2)

    async with live.stream("GET", "/api/stream") as response:
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("text/event-stream")
        talking: asyncio.Task[None] = asyncio.create_task(chatter())
        try:
            async with aclosing(follow_sse(response)) as messages:
                first: list[SseMessage] = [await asyncio.wait_for(anext(messages), 15.0) for _ in range(3)]
        finally:
            talking.cancel()
    assert [item.kind for item in first] == ["run_started"] * 3
    assert {item.data.run_id for item in first} == {run.id}


async def test_an_open_stream_ends_with_the_server_instead_of_holding_up_its_shutdown(
    runtime: Runtime, manager: RunManager
) -> None:
    server, serving, address = await _start_server(runtime, manager)
    try:
        async with (
            httpx.AsyncClient(
                base_url=address, headers={"Authorization": f"Bearer {TOKEN}"}, timeout=SSE_TIMEOUT_S
            ) as http,
            http.stream("GET", "/api/stream") as response,
        ):
            assert response.status_code == 200
            server.should_exit = True
            await asyncio.wait_for(asyncio.shield(serving), STOP_WITHIN_S)
            # The stream was completed, not cut off: reading it to the end raises nothing.
            assert await response.aread() == b""
    finally:
        server.should_exit = True
        await serving


# ------------------------------------------------------------------------------------------ runs with a model


async def _answer_until_the_end(client: httpx.AsyncClient, run_id: str, approve: bool) -> RunRecord:
    """Give every approval the run asks for the same answer until the run is over."""
    deadline: float = time.monotonic() + RUN_TIMEOUT_S
    while True:
        detail: httpx.Response = await client.get(f"/api/runs/{run_id}")
        assert detail.status_code == 200, detail.text
        if not detail.json()["active"]:
            return RunRecord.model_validate(detail.json()["run"])
        pending: httpx.Response = await client.get("/api/approvals", params={"status": "pending"})
        for item in pending.json():
            if item["run_id"] == run_id:
                decided: httpx.Response = await client.post(
                    f"/api/approvals/{item['id']}/decision", json={"approve": approve, "note": ""}
                )
                assert decided.status_code == 200, decided.text
        assert time.monotonic() < deadline, f"run {run_id} did not finish"
        await asyncio.sleep(POLL_S)


async def _approve_until_the_end(client: httpx.AsyncClient, run_id: str) -> RunRecord:
    return await _answer_until_the_end(client, run_id, True)


async def _deny_until_the_end(client: httpx.AsyncClient, run_id: str) -> RunRecord:
    return await _answer_until_the_end(client, run_id, False)


def _shell_results(events: list[EventRecord]) -> list[str]:
    """What the model read back from each shell call of a run."""
    return [
        str(event.data["text"]) for event in events if event.kind == "tool_result" and event.data["tool"] == "shell"
    ]


def _text_of(path: Path) -> str | None:
    """What the file holds; ``None`` when there is no such file."""
    return path.read_text(encoding="utf-8") if path.exists() else None


def _secret_file(workspace: Path) -> tuple[Path, str]:
    token: str = f"MARKER-{secrets.token_hex(6)}"
    path: Path = workspace / f"delegate-{secrets.token_hex(3)}.txt"
    path.write_text(token, encoding="utf-8")
    return path, token


@requires_model
@pytest.mark.live
async def test_a_run_started_through_the_api_finishes_and_is_replayed_by_its_stream(client: httpx.AsyncClient) -> None:
    started: RunRecord = await _start(client, "Call the finish tool right away with the answer pong and success true.")
    assert (started.source, started.mode, started.model, started.parent_run_id) == ("ui", "task", MODEL_REF, None)
    finished: RunRecord = await _run_to_the_end(client, started.id)
    assert finished.status == "done", finished
    assert "pong" in finished.result.lower()
    events: list[EventRecord] = await _events(client, started.id)
    assert [events[0].kind, events[-1].kind] == ["run_queued", "run_finished"]
    assert [event.id for event in events] == sorted({event.id for event in events})
    later: httpx.Response = await client.get(f"/api/runs/{started.id}/events", params={"after": events[1].id})
    assert [item["id"] for item in later.json()] == [event.id for event in events[2:]]
    streamed: httpx.Response = await client.get(f"/api/runs/{started.id}/stream")
    assert [item.id for item in parse_sse(streamed.text)] == [event.id for event in events]
    listed: httpx.Response = await client.get("/api/runs", params={"thread_id": started.thread_id})
    assert [run["id"] for run in listed.json()] == [started.id]
    assert started.id in [run["id"] for run in (await client.get("/api/runs", params={"limit": 200})).json()]


@requires_model
@pytest.mark.live
async def test_a_run_keeps_working_after_the_page_that_followed_it_is_left(
    runtime: Runtime, manager: RunManager
) -> None:
    server, serving, address = await _start_server(runtime, manager)
    try:
        async with httpx.AsyncClient(
            base_url=address, headers={"Authorization": f"Bearer {TOKEN}"}, timeout=SSE_TIMEOUT_S
        ) as http:
            started: httpx.Response = await http.post(
                "/api/runs",
                json={
                    "goal": "Think about the number 7 for a moment, then call the finish tool with the answer pong.",
                    "approvals": "yolo",
                },
            )
            assert started.status_code == 201, started.text
            run_id: str = started.json()["id"]
            # The page follows the run, reads its first event and is left (the stream is closed).
            async with http.stream("GET", f"/api/runs/{run_id}/stream") as response:
                assert response.status_code == 200
                async for line in response.aiter_lines():
                    if line.startswith("data:"):
                        break
            detail: httpx.Response = await http.get(f"/api/runs/{run_id}")
            assert detail.json()["run"]["status"] in {"queued", "running", "done"}
            async with asyncio.timeout(RUN_TIMEOUT_S):
                finished: RunRecord = await manager.wait(run_id)
            assert finished.status == "done", finished
            assert "pong" in finished.result.lower()
    finally:
        server.should_exit = True
        await serving


@requires_model
@pytest.mark.live
async def test_a_shell_command_waits_for_the_approval_of_the_user_and_runs_once_it_is_given(
    client: httpx.AsyncClient, runtime: Runtime
) -> None:
    token: str = f"MARKER-{secrets.token_hex(6)}"
    path: Path = runtime.workspace / f"approved-{token}.txt"
    goal: str = f"Run exactly this shell command: touch {path}  Then call finish with the answer done."
    run: RunRecord = await _start(client, goal)
    approval: ApprovalRecord = await _pending_approval(client, run.id)
    assert approval.tool == "shell"
    assert token in f"{approval.summary} {approval.args}"
    assert _text_of(path) is None
    note: str = f"NOTE-{secrets.token_hex(4)}"
    decided: httpx.Response = await client.post(
        f"/api/approvals/{approval.id}/decision", json={"approve": True, "note": note}
    )
    assert decided.status_code == 200, decided.text
    assert (decided.json()["status"], decided.json()["decided_by"]) == ("approved", "ui")
    again: httpx.Response = await client.post(
        f"/api/approvals/{approval.id}/decision", json={"approve": False, "note": ""}
    )
    assert again.status_code == 409
    finished: RunRecord = await _approve_until_the_end(client, run.id)
    assert finished.status == "done", finished
    assert _text_of(path) == ""
    events: list[EventRecord] = await _events(client, run.id)
    kinds: list[str] = [event.kind for event in events]
    assert kinds.index("approval_requested") < kinds.index("approval_decided") < kinds.index("run_finished")
    assert any(note in text for text in _shell_results(events)), "the model never read the note of the approval"


@requires_model
@pytest.mark.live
async def test_a_shell_command_the_user_denies_does_not_run(client: httpx.AsyncClient, runtime: Runtime) -> None:
    path: Path = runtime.workspace / f"denied-{secrets.token_hex(3)}.txt"
    goal: str = f"Run exactly this shell command: touch {path}  If it is not allowed, call finish and say so."
    run: RunRecord = await _start(client, goal)
    approval: ApprovalRecord = await _pending_approval(client, run.id)
    denied: httpx.Response = await client.post(
        f"/api/approvals/{approval.id}/decision", json={"approve": False, "note": "not now"}
    )
    assert denied.status_code == 200, denied.text
    assert (denied.json()["status"], denied.json()["note"]) == ("denied", "not now")
    await _deny_until_the_end(client, run.id)
    assert _text_of(path) is None
    assert any("not now" in text for text in _shell_results(await _events(client, run.id)))


@requires_model
@pytest.mark.live
async def test_a_question_of_the_run_is_answered_with_the_note_of_the_decision(client: httpx.AsyncClient) -> None:
    codeword: str = f"MARKER-{secrets.token_hex(6)}"
    goal: str = (
        "Call ask_user with the question 'Which codeword should I use?'. "
        "Then call finish with the exact codeword the user gave as the answer."
    )
    run: RunRecord = await _start(client, goal)
    question: ApprovalRecord = await _pending_approval(client, run.id)
    assert (question.tool, question.severity) == ("ask_user", "question")
    answered: httpx.Response = await client.post(
        f"/api/approvals/{question.id}/decision", json={"approve": True, "note": codeword}
    )
    assert answered.status_code == 200, answered.text
    finished: RunRecord = await _run_to_the_end(client, run.id)
    assert finished.status == "done", finished
    assert codeword in finished.result


@requires_model
@pytest.mark.live
async def test_a_run_started_in_yolo_mode_is_never_asked_what_the_cautious_preset_would_ask(
    client: httpx.AsyncClient, runtime: Runtime
) -> None:
    path: Path = runtime.workspace / f"yolo-{secrets.token_hex(3)}.txt"
    goal: str = f"Run exactly this shell command: touch {path}  Then call finish with the answer done."
    answer: httpx.Response = await client.post("/api/runs", json={"goal": goal, "approvals": "yolo"})
    assert answer.status_code == 201, answer.text
    started: RunRecord = RunRecord.model_validate(answer.json())
    assert (started.approvals, started.goal_mode) == ("yolo", False)
    finished: RunRecord = await _run_to_the_end(client, started.id)
    assert finished.status == "done", finished
    assert _text_of(path) is not None
    events: list[EventRecord] = await _events(client, started.id)
    assert "approval_requested" not in [event.kind for event in events]
    assert {event.data["decided_by"] for event in events if event.kind == "approval_auto"} == {"yolo"}


@requires_model
@pytest.mark.live
async def test_a_run_waiting_for_an_approval_can_be_steered_and_cancelled(
    client: httpx.AsyncClient, runtime: Runtime
) -> None:
    path: Path = runtime.workspace / f"cancelled-{secrets.token_hex(3)}.txt"
    run: RunRecord = await _start(client, f"Run exactly this shell command: touch {path}  Then call finish.")
    approval: ApprovalRecord = await _pending_approval(client, run.id)
    steered: httpx.Response = await client.post(f"/api/runs/{run.id}/steer", json={"text": "Take your time."})
    assert steered.status_code == 202, steered.text
    assert steered.json()["active"] is True
    cancelled: httpx.Response = await client.post(f"/api/runs/{run.id}/cancel")
    assert cancelled.status_code == 202, cancelled.text
    finished: RunRecord = await _run_to_the_end(client, run.id)
    assert finished.status == "cancelled", finished
    pending: httpx.Response = await client.get("/api/approvals", params={"status": "pending"})
    assert approval.id not in [item["id"] for item in pending.json()]
    closed: httpx.Response = await client.get("/api/approvals", params={"status": "expired"})
    assert [(item["decided_by"], item["note"]) for item in closed.json() if item["id"] == approval.id] == [
        ("system", f"run {run.id} stopped waiting")
    ]
    decided: list[object] = [
        event.data["approval_id"] for event in await _events(client, run.id) if event.kind == "approval_decided"
    ]
    assert decided == [approval.id]
    assert _text_of(path) is None
    assert (await client.post(f"/api/runs/{run.id}/steer", json={"text": "too late"})).status_code == 409
    assert (await client.post(f"/api/runs/{run.id}/cancel")).status_code == 409


# ------------------------------------------------------------------------------------------ chats


async def _chat(runtime: Runtime, title: str, runs: int) -> tuple[ThreadRecord, list[RunRecord]]:
    """A thread with ``runs`` top-level runs that only exist in the store, oldest first."""
    thread: ThreadRecord = await runtime.store.create_thread(title, "test")
    made: list[RunRecord] = []
    for _ in range(runs):
        made.append(
            await runtime.store.create_run(
                thread.id,
                f"{title} goal",
                "chat",
                MODEL_REF,
                None,
                "test",
                "medium",
                "standard",
                None,
                "ask",
                False,
                None,
                None,
                None,
            )
        )
    return thread, made


def _said(text: str) -> Message:
    return Message(
        role="user",
        content=text,
        images=(),
        tool_calls=(),
        tool_call_id=None,
        tool_name=None,
        thinking=None,
        raw=(),
        raw_origin=None,
    )


async def _chat_ids(client: httpx.AsyncClient) -> list[str]:
    answer: httpx.Response = await client.get("/api/threads", params={"limit": 500})
    assert answer.status_code == 200, answer.text
    return [item["id"] for item in answer.json()]


async def test_the_chats_are_the_threads_a_person_started_runs_on_with_their_latest_run(
    client: httpx.AsyncClient, runtime: Runtime
) -> None:
    first, first_runs = await _chat(runtime, "first chat", 2)
    second, (second_run,) = await _chat(runtime, "second chat", 1)
    delegated: RunRecord = await _stored_run(runtime, second_run.id)
    answer: httpx.Response = await client.get("/api/threads", params={"limit": 500})
    assert answer.status_code == 200, answer.text
    chats: dict[str, ThreadSummary] = {item["id"]: ThreadSummary.model_validate(item) for item in answer.json()}
    assert delegated.thread_id not in chats, "a thread holding only a delegated run is not a chat"
    assert (chats[first.id].run_count, chats[first.id].latest_run_id) == (2, first_runs[-1].id)
    assert (chats[first.id].title, chats[first.id].latest_status, chats[first.id].latest_mode) == (
        "first chat",
        "queued",
        "chat",
    )
    assert chats[first.id].updated_at == first_runs[-1].created_at, "the start of the latest run counts as activity"
    assert (chats[second.id].run_count, chats[second.id].latest_run_id) == (1, second_run.id)
    order: list[str] = [item["id"] for item in answer.json()]
    assert order.index(second.id) < order.index(first.id), "the latest activity comes first"
    one: httpx.Response = await client.get("/api/threads", params={"limit": 1})
    assert [item["id"] for item in one.json()] == [order[0]]
    assert (await client.get("/api/threads", params={"limit": 0})).status_code == 422


async def test_deleting_chats_takes_their_runs_delegations_events_and_approvals_along(
    client: httpx.AsyncClient, runtime: Runtime
) -> None:
    doomed, (top,) = await _chat(runtime, "doomed chat", 1)
    child: RunRecord = await _stored_run(runtime, top.id)
    grandchild: RunRecord = await _stored_run(runtime, child.id)
    # The thread of the grandchild also holds a run of its own, so it stays; the thread of the child does not.
    own: RunRecord = await runtime.store.create_run(
        grandchild.thread_id,
        "own goal",
        "task",
        MODEL_REF,
        None,
        "test",
        "medium",
        "standard",
        None,
        "ask",
        False,
        None,
        None,
        None,
    )
    kept_chat, (kept,) = await _chat(runtime, "kept chat", 1)
    await runtime.store.append_messages(doomed.id, [_said("hello")])
    await _tell(
        runtime,
        [
            (top.id, "run_queued", _queued(None)),
            (child.id, "run_queued", _queued(top.id)),
            (grandchild.id, "run_queued", _queued(child.id)),
            (kept.id, "run_queued", _queued(None)),
        ],
    )
    await runtime.store.create_approval(grandchild.id, "shell", "touch x", "write", {"command": "touch x"})
    await runtime.store.create_approval(kept.id, "shell", "touch y", "write", {"command": "touch y"})
    # A queue on the child's thread (as the API took one before item 170) goes with the thread; the grandchild's
    # thread, a chat of its own, keeps its queue.
    for thread_id in (child.thread_id, grandchild.thread_id):
        await runtime.store.add_queued(
            thread_id, "next", [], "task", None, None, "standard", None, None, None, None, "test"
        )
        await runtime.store.hold_queue(thread_id)
    # Runs this daemon does not drive, as a dotx command's, keep their chat while they work (item 179).
    refused: httpx.Response = await client.request("DELETE", "/api/threads", json={"ids": [doomed.id]})
    assert refused.status_code == 409, refused.text
    assert refused.json()["context"] == {
        "thread_ids": [doomed.id],
        "run_ids": sorted([top.id, child.id, grandchild.id]),
    }
    for over in (top, child, grandchild):
        await runtime.store.finish_run(over.id, "done", "Done.")

    answer: httpx.Response = await client.request(
        "DELETE", "/api/threads", json={"ids": [doomed.id, "thr_nothing", doomed.id]}
    )
    assert answer.status_code == 200, answer.text
    assert answer.json() == {"deleted": [doomed.id], "missing": ["thr_nothing"]}

    for run_id in (top.id, child.id, grandchild.id):
        assert (await client.get(f"/api/runs/{run_id}")).status_code == 404, run_id
        assert await runtime.store.list_events(run_id, 0, 10) == []
    assert (await client.get(f"/api/runs/{own.id}")).status_code == 200
    assert (await client.get(f"/api/runs/{kept.id}")).status_code == 200
    pending: httpx.Response = await client.get("/api/approvals", params={"status": "pending"})
    asked: list[str] = [item["run_id"] for item in pending.json()]
    assert grandchild.id not in asked and kept.id in asked
    with pytest.raises(NotFoundError):
        await runtime.store.get_thread(doomed.id)
    with pytest.raises(NotFoundError):
        await runtime.store.get_thread(child.thread_id)
    assert (await runtime.store.get_thread(grandchild.thread_id)).id == grandchild.thread_id
    assert await runtime.store.list_queued(child.thread_id) == []
    assert not await runtime.store.queue_held(child.thread_id)
    assert [message.text for message in await runtime.store.list_queued(grandchild.thread_id)] == ["next"]
    assert await runtime.store.queue_held(grandchild.thread_id)
    assert await runtime.store.load_messages(doomed.id, 0) == []
    chats: list[str] = await _chat_ids(client)
    assert doomed.id not in chats and kept_chat.id in chats and grandchild.thread_id in chats
    assert (await client.request("DELETE", "/api/threads", json={"ids": []})).status_code == 422
    assert (await client.request("DELETE", "/api/threads", json={})).status_code == 422


async def test_a_chat_is_named_and_pinned_from_its_menu(client: httpx.AsyncClient, runtime: Runtime) -> None:
    older, _ = await _chat(runtime, "older chat", 1)
    newer, _ = await _chat(runtime, "newer chat", 1)
    assert (await client.patch(f"/api/threads/{older.id}", json={"pinned": True})).status_code == 204
    assert (await client.patch(f"/api/threads/{older.id}", json={"title": "  Trip notes  "})).status_code == 204
    listed: dict[str, ThreadSummary] = {
        item["id"]: ThreadSummary.model_validate(item)
        for item in (await client.get("/api/threads", params={"limit": 500})).json()
    }
    assert (listed[older.id].title, listed[newer.id].pinned_at) == ("Trip notes", None)
    assert listed[older.id].pinned_at is not None
    order: list[str] = await _chat_ids(client)
    assert order.index(older.id) < order.index(newer.id), "a pinned chat comes before the latest activity"
    one: httpx.Response = await client.get("/api/threads", params={"limit": 1})
    assert [item["id"] for item in one.json()] == [older.id], "a pinned chat is listed however old"
    assert (await client.patch(f"/api/threads/{older.id}", json={"pinned": False})).status_code == 204
    order = await _chat_ids(client)
    assert order.index(newer.id) < order.index(older.id), "unpinned, it goes back to its activity's place"
    assert (await client.patch("/api/threads/thr_nothing", json={"pinned": True})).status_code == 404
    for wrong in ({}, {"title": "   "}, {"title": "x" * 121}, {"pinned": "soon"}, {"pinned": "yes"}, {"pinned": 1}):
        assert (await client.patch(f"/api/threads/{older.id}", json=wrong)).status_code == 422, wrong


async def _unread(client: httpx.AsyncClient, thread_id: str) -> bool:
    listed: list[ThreadSummary] = [
        ThreadSummary.model_validate(item) for item in (await client.get("/api/threads", params={"limit": 500})).json()
    ]
    return next(chat.unread for chat in listed if chat.id == thread_id)


async def test_a_chat_is_marked_read_and_unread_from_its_menu(client: httpx.AsyncClient, runtime: Runtime) -> None:
    chat, runs = await _chat(runtime, "ended chat", 1)
    assert not await _unread(client, chat.id), "a run under way leaves its chat read"
    await runtime.store.finish_run(runs[0].id, "done", "Done.")
    assert await _unread(client, chat.id), "a run that ended in a chat not seen since makes it unread"
    assert (await client.patch(f"/api/threads/{chat.id}", json={"seen": True})).status_code == 204
    assert not await _unread(client, chat.id)
    assert (await client.patch(f"/api/threads/{chat.id}", json={"seen": False})).status_code == 204
    assert await _unread(client, chat.id)
    assert (await client.patch("/api/threads/thr_nothing", json={"seen": True})).status_code == 404
    for wrong in ({"seen": None}, {"seen": "yes"}, {"seen": 1}):
        assert (await client.patch(f"/api/threads/{chat.id}", json=wrong)).status_code == 422, wrong


@requires_model
@pytest.mark.live
async def test_a_chat_with_a_run_under_way_is_kept_until_the_run_is_over(
    client: httpx.AsyncClient, runtime: Runtime
) -> None:
    path: Path = runtime.workspace / f"busy-{secrets.token_hex(3)}.txt"
    goal: str = f"Run exactly this shell command: touch {path}  If it is not allowed, call finish and say so."
    run: RunRecord = await _start(client, goal)
    await _pending_approval(client, run.id)
    refused: httpx.Response = await client.request("DELETE", "/api/threads", json={"ids": [run.thread_id]})
    assert refused.status_code == 409, refused.text
    assert refused.json()["context"] == {"thread_ids": [run.thread_id]}
    assert (await client.get(f"/api/runs/{run.id}")).status_code == 200
    assert run.thread_id in await _chat_ids(client)
    await _deny_until_the_end(client, run.id)
    deleted: httpx.Response = await client.request("DELETE", "/api/threads", json={"ids": [run.thread_id]})
    assert deleted.status_code == 200, deleted.text
    assert deleted.json() == {"deleted": [run.thread_id], "missing": []}
    assert (await client.get(f"/api/runs/{run.id}")).status_code == 404
    assert run.thread_id not in await _chat_ids(client)
    assert _text_of(path) is None


# ------------------------------------------------------------------------------------------ live streams


async def _follow_approving(live: httpx.AsyncClient, run_id: str) -> list[SseMessage]:
    """Follow the stream of a run to its end, approving whatever it asks for; every message that came."""
    seen: list[SseMessage] = []
    async with asyncio.timeout(RUN_TIMEOUT_S), live.stream("GET", f"/api/runs/{run_id}/stream") as response:
        assert response.status_code == 200, await response.aread()
        assert response.headers["content-type"].startswith("text/event-stream")
        async with aclosing(follow_sse(response)) as messages:
            async for message in messages:
                seen.append(message)
                if message.kind == "approval_requested":
                    decided: httpx.Response = await live.post(
                        f"/api/approvals/{message.data.data['approval_id']}/decision",
                        json={"approve": True, "note": ""},
                    )
                    assert decided.status_code == 200, decided.text
    return seen


@requires_model
@pytest.mark.live
async def test_a_live_stream_shows_a_run_as_it_happens_and_ends_with_it(
    live: httpx.AsyncClient, runtime: Runtime
) -> None:
    path: Path = runtime.workspace / f"live-{secrets.token_hex(3)}.txt"
    goal: str = f"Run exactly this shell command: touch {path}  Then call finish with the answer done."
    run: RunRecord = await _start(live, goal)
    seen: list[SseMessage] = await _follow_approving(live, run.id)
    kinds: list[str] = [item.kind for item in seen]
    assert kinds[0] == "run_queued"
    assert kinds[-1] == "run_finished"
    assert kinds.index("approval_requested") < kinds.index("approval_decided") < kinds.index("run_finished")
    identifiers: list[int] = [item.id for item in seen if item.id is not None]
    assert identifiers == sorted(set(identifiers))
    assert _text_of(path) is not None


@requires_model
@pytest.mark.live
async def test_a_live_stream_carries_the_runs_the_run_delegated_to(live: httpx.AsyncClient, runtime: Runtime) -> None:
    path, token = _secret_file(runtime.workspace)
    goal: str = (
        f"Call the delegate tool exactly once, with mode task. Its goal must tell the sub-agent to read the file {path} "
        "and report the exact text in it. When it has answered, call finish with the text it reported."
    )
    run: RunRecord = await _start(live, goal)
    seen: list[SseMessage] = await _follow_approving(live, run.id)
    queued: list[SseMessage] = [item for item in seen if item.kind == "run_queued"]
    assert len(queued) == 2, [item.kind for item in seen]
    assert queued[0].data.run_id == run.id
    child_id: str = queued[1].data.run_id
    assert queued[1].data.data["parent_run_id"] == run.id
    finishes: list[str] = [item.data.run_id for item in seen if item.kind == "run_finished"]
    assert finishes == [child_id, run.id]
    assert seen[-1].data.run_id == run.id
    finished: RunRecord = await _run_to_the_end(live, run.id)
    assert finished.status == "done", finished
    assert token in finished.result


@pytest.mark.browser
async def test_the_web_ui_signs_in_from_the_address_dotx_open_builds_and_refuses_a_wrong_token(
    live: httpx.AsyncClient, runtime: Runtime
) -> None:
    """The first thing a person does with the UI: open the address with the token, or paste a token into the form."""
    port: int | None = live.base_url.port
    assert port is not None
    shell: str = ".shell"
    field: str = "input[aria-label='Access token']"
    async with async_playwright() as playwright:
        browser: Browser = await playwright.chromium.launch(channel=runtime.config.browser.channel, headless=True)
        try:
            opened: Page = await browser.new_page()
            await opened.goto(ui_url(live.base_url.host, port, TOKEN))
            await expect(opened.locator(shell)).to_be_visible()
            assert TOKEN not in opened.url, "the token must leave the address bar"
            await opened.reload()
            await expect(opened.locator(shell)).to_be_visible()

            by_hand: Page = await browser.new_page()
            await by_hand.goto(f"http://{live.base_url.host}:{port}/")
            await by_hand.fill(field, "not-the-token")
            await by_hand.press(field, "Enter")
            await expect(by_hand.get_by_text("The daemon did not accept that access token.")).to_be_visible()
            await by_hand.fill(field, TOKEN)
            await by_hand.press(field, "Enter")
            await expect(by_hand.locator(shell)).to_be_visible()
        finally:
            await browser.close()


@pytest.mark.browser
@requires_model
@pytest.mark.live
async def test_a_file_attached_in_the_web_ui_reaches_the_model(live: httpx.AsyncClient, runtime: Runtime) -> None:
    """Attach a file in the composer, see it uploaded, start the run, and get an answer only the file holds."""
    port: int | None = live.base_url.port
    assert port is not None
    code: str = f"heron-{secrets.token_hex(3)}"
    attached: FilePayload = {
        "name": "clue.txt",
        "mimeType": "text/plain",
        "buffer": f"The code word is {code}.\n".encode(),
    }
    async with async_playwright() as playwright:
        browser: Browser = await playwright.chromium.launch(channel=runtime.config.browser.channel, headless=True)
        try:
            page: Page = await browser.new_page()
            await page.goto(ui_url(live.base_url.host, port, TOKEN))
            await page.locator(".composer input[type=file]").set_input_files(attached)
            await expect(page.locator(".composer .attach-chip.ready")).to_contain_text("clue.txt")
            goal: str = "textarea[aria-label='Goal']"
            await page.fill(goal, "What is the code word in the attached file? Reply with the code word only.")
            await page.press(goal, "Enter")
            await expect(page.locator(".attach-list.sent")).to_contain_text("clue.txt")
            await expect(page.locator("article.answer")).to_contain_text(code, timeout=RUN_TIMEOUT_S * 1000)
        finally:
            await browser.close()


async def test_a_fault_and_a_path_that_does_not_exist_are_answered_as_json(live: httpx.AsyncClient) -> None:
    """Item 178: an error that is not one of Dot's own (here the themes folder cannot be listed) and an answer of the web
    framework itself (a path that does not exist) come in the shape of every other error of the API, which a page reads;
    the fault is logged with its traceback. Before, the fault came as the plain text "Internal Server Error"."""
    folder: Path = themes_dir()
    await asyncio.to_thread(folder.mkdir, parents=True, exist_ok=True)
    await asyncio.to_thread(folder.chmod, 0)
    try:
        with capture_logs() as logs:
            failed: httpx.Response = await live.get("/api/themes")
        assert failed.status_code == 500, failed.text
        body: JsonObject = failed.json()
        assert set(body) == {"message", "context"}, body
        error = body["context"]
        assert isinstance(error, dict) and str(error["error"]).startswith("PermissionError"), body
        crashed: list[JsonObject] = [entry for entry in logs if entry["event"] == "request_crashed"]
        assert [entry["path"] for entry in crashed] == ["/api/themes"], logs
        assert crashed[0]["exc_info"] is not None
    finally:
        await asyncio.to_thread(folder.chmod, 0o755)
    missing: httpx.Response = await live.get("/api/no-such-path-178")
    assert missing.status_code == 404, missing.text
    assert set(missing.json()) == {"message", "context"}, missing.text
