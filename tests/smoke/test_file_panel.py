"""Files in the panel beside the conversation (item 195), on the pages of a daemon in Chrome. A file a message carries
opens beside the conversation when its chip is clicked: a PDF's pages as pictures the daemon draws; a workbook, a
document and a presentation first as macOS's Quick Look draws them, in a sandboxed frame (item 268, the real
``qlmanage``), then, a click away, as a grid with a tab for each sheet and what a cell holds, a document's headings,
list and table, and a presentation's slides with their charts. A file the answer names by its path opens the same
way: a CSV file as a grid, which opens on the Mac, an image, and a Markdown file formatted, its HTML never drawn,
with its source a click away (item 269); a path the panel does not show says why. The side panel's Run tab lists
no files, which open from the conversation (item 544). A long text pasted into the
conversation's box goes as a file, which opens beside the conversation from its chip (item 269). Closing the file
gives the panel back as it was and the focus to what opened it. Wide and on a phone, light and dark. The daemon's
routes refuse a file outside the home folder and a script to open, and serve a PDF's page as a PNG.
No network and no model service is used: the scripted model plays the runs of a daemon served by a real server, the
home folder the daemon shows files of is a scratch one, and a stand-in takes the place of Launch Services: a test
never opens an application.
"""

import asyncio
import io
import re
import time
from collections.abc import AsyncIterator
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Literal

import docx
import openpyxl
import pptx
import pytest
import pytest_asyncio
from PIL import Image, ImageDraw
from playwright.async_api import APIResponse, Browser, FrameLocator, Locator, Page, async_playwright, expect
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from dotagent.agent.attachments import StoredAttachment, store_attachment
from dotagent.agent.manager import RunManager
from dotagent.config import load_default_config
from dotagent.evals.runner import SCRIPTED_MODEL, eval_config, eval_environment, scratch_home
from dotagent.evals.scripted import ScriptedProvider
from dotagent.evals.types import ScriptedCall, ScriptedTurn
from dotagent.llm.router import parse_model_ref
from dotagent.office import OFFICE_MEDIA_TYPES
from dotagent.previews import quicklook
from dotagent.runtime import Runtime, close_runtime, open_runtime
from dotagent.server import routes_files
from dotagent.server.app import create_app
from dotagent.server.context import ServerContext
from dotagent.server.daemon import DotServer, bind_listener, serving_config
from dotagent.server.links import ui_url
from dotagent.store.models import RunRecord
from dotagent.tools.meta import FINISH_TOOL
from dotagent.tools.shell import read_login_env

pytestmark = pytest.mark.asyncio(loop_scope="module")

TOKEN: str = "file-panel-test-token"
WAIT_S: float = 30.0
ANSWER: str = (
    "The rows are in `~/Documents/rows.csv`, the plan in `~/Documents/todo.md`, the chart in"
    " `~/Documents/chart.png` and the script in `~/Documents/total.py`; the notes in `~/.private/notes.txt` stay where"
    " they are."
)
# A source file, which the panel shows as code: its lines numbered and its words coloured by kind (item 624).
SCRIPT: str = 'def total(rows):\n    # The sum of the amounts.\n    return sum(row["amount"] for row in rows) // 1\n'

# Markup in a text file, which the panel shows as it is written.
TODO: str = "# To do\n\nPay the <b>rent</b> first.\n"
# What each run writes in the workspace before it answers.
SUMMARY: str = "# Summary\n\nAll paid.\n"
# A long text pasted into a conversation's box (item 269).
PASTED: str = "# Pasted plan\n\n" + "".join(f"- Point {number}\n" for number in range(1, 41))
_PROVIDER: str = parse_model_ref(SCRIPTED_MODEL).provider
_SHOTS: tuple[tuple[int, Literal["light", "dark"]], ...] = (
    (1280, "light"),
    (1280, "dark"),
    (390, "light"),
    (390, "dark"),
)
_FINISH: ScriptedTurn = ScriptedTurn(
    text="", calls=(ScriptedCall(tool=FINISH_TOOL, arguments={"answer": ANSWER, "success": True}),)
)


def _turns(workspace: Path) -> list[ScriptedTurn]:
    """Each run writes the summary in the workspace, then answers; the run a message sent from the box starts (item
    406) does the same."""
    write: ScriptedTurn = ScriptedTurn(
        text="",
        calls=(ScriptedCall(tool="write_file", arguments={"path": str(_summary(workspace)), "content": SUMMARY}),),
    )
    return [write, _FINISH] * (2 * len(_SHOTS))


def _summary(workspace: Path) -> Path:
    return workspace / "made" / "summary.md"


def _pdf() -> bytes:
    """A PDF of two pages of US Letter, each saying its number."""
    pages: list[Image.Image] = [Image.new("RGB", (612, 792), (255, 255, 255)) for _ in range(2)]
    for number, picture in enumerate(pages, 1):
        ImageDraw.Draw(picture).text((72, 72), f"Page {number} of the report", fill=(0, 0, 0))
    out: io.BytesIO = io.BytesIO()
    pages[0].save(out, format="PDF", save_all=True, append_images=pages[1:], resolution=72.0)
    return out.getvalue()


def _workbook() -> bytes:
    """A budget with a formula no program that calculates has worked out, and a hidden sheet of notes."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Budget"
    sheet.append(["Item", "Amount"])
    sheet.append(["Rent", 1200])
    sheet.append(["Food", 300])
    sheet["B4"] = "=SUM(B2:B3)"
    notes = workbook.create_sheet("Notes")
    notes["A1"] = "Paid on the first"
    notes.sheet_state = "hidden"
    out: io.BytesIO = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _document() -> bytes:
    document = docx.Document()
    document.add_heading("Budget 2027", level=1)
    document.add_paragraph("The plan for the year.")
    document.add_paragraph("Rent", style="List Bullet")
    document.add_paragraph("Food", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    for row, cells in zip(table.rows, (("Item", "Amount"), ("Rent", "1200")), strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    out: io.BytesIO = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _presentation() -> bytes:
    """A slide with a title, a point and a chart of two quarters."""
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    assert slide.shapes.title is not None
    slide.shapes.title.text = "Sales"
    slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1)).text_frame.text = "Up in the spring"
    chart = CategoryChartData()
    chart.categories = ["Q1", "Q2"]
    chart.add_series("2027", (1.0, 2.5))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(3), chart)
    out: io.BytesIO = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def _picture() -> bytes:
    out: io.BytesIO = io.BytesIO()
    Image.new("RGB", (320, 200), (40, 120, 200)).save(out, format="PNG")
    return out.getvalue()


def _user_home(root: Path) -> Path:
    """A home folder with a table in Documents, notes in a hidden folder and a script."""
    home: Path = root / "user"
    (home / "Documents").mkdir(parents=True)
    (home / ".private").mkdir()
    (home / "Documents" / "rows.csv").write_text("Item;Amount\nRent;1200,50\nFood;300\n", encoding="utf-8")
    (home / "Documents" / "todo.md").write_text(TODO, encoding="utf-8")
    (home / "Documents" / "chart.png").write_bytes(_picture())
    (home / "Documents" / "total.py").write_text(SCRIPT, encoding="utf-8")
    (home / ".private" / "notes.txt").write_text("Private notes.\n", encoding="utf-8")
    (home / "Documents" / "tidy.sh").write_text("#!/bin/sh\necho tidy\n", encoding="utf-8")
    return home


@pytest_asyncio.fixture(scope="module", loop_scope="module")
async def login_env() -> dict[str, str]:
    return await read_login_env(load_default_config().shell.login_env_timeout_s, ())


@asynccontextmanager
async def _daemon(login_env: dict[str, str], user_home: Path) -> AsyncIterator[tuple[Runtime, RunManager, str]]:
    """The daemon in the home of the environment, whose model writes the summary and answers ``ANSWER`` once a size,
    and which shows the files of ``user_home``, served at the address it yields."""
    runtime: Runtime = await open_runtime(eval_config(load_default_config()), login_env)
    try:
        await runtime.router.register(_PROVIDER, ScriptedProvider(_PROVIDER, _turns(runtime.workspace)))
        manager: RunManager = RunManager(runtime)
        stopping: asyncio.Event = asyncio.Event()
        listener = bind_listener("127.0.0.1", 0)
        port: int = listener.getsockname()[1]
        context: ServerContext = ServerContext(
            runtime=runtime,
            manager=manager,
            scheduler=None,
            token=TOKEN,
            host="127.0.0.1",
            port=port,
            started_at=time.time(),
            stopping=stopping,
        )
        # The file routes take the home folder of the daemon's user when the app is made: here the scratch one.
        with pytest.MonkeyPatch.context() as patch:
            patch.setenv("HOME", str(user_home))
            server: DotServer = DotServer(serving_config(create_app(context)), stopping)
        serving: asyncio.Task[None] = asyncio.create_task(server.serve(sockets=[listener]))
        while not server.started:
            assert not serving.done(), "the server stopped while starting"
            await asyncio.sleep(0.05)
        try:
            yield runtime, manager, f"http://127.0.0.1:{port}"
        finally:
            server.should_exit = True
            stopping.set()
            await serving
            await manager.shutdown()
    finally:
        await close_runtime(runtime)


async def _close(page: Page, view: Locator, opener: Locator) -> None:
    """Close the file with its button: the panel lets it go and the focus goes back to what opened it."""
    await view.get_by_role("button", name="Close the file").click()
    await expect(view).to_have_count(0)
    await expect(opener).to_be_focused()
    await expect(page.locator(".run-layout")).not_to_have_attribute("data-file", "on")


async def _one_size(
    browser: Browser,
    runtime: Runtime,
    manager: RunManager,
    address: str,
    files: tuple[StoredAttachment, ...],
    user_home: Path,
    opened: list[Path],
    size: tuple[int, Literal["light", "dark"]],
    shots: Path,
) -> None:
    width, scheme = size
    page: Page = await browser.new_page(
        viewport={"width": width, "height": 900}, color_scheme=scheme, reduced_motion="reduce"
    )
    await page.goto(ui_url("127.0.0.1", int(address.rpartition(":")[2]), TOKEN))
    run: RunRecord = await manager.start(
        f"Look at the files ({width}, {scheme})",
        "task",
        SCRIPTED_MODEL,
        None,
        "test",
        None,
        None,
        "standard",
        None,
        "ask",
        False,
        files,
    )
    async with asyncio.timeout(WAIT_S):
        assert (await manager.wait(run.id)).status == "done"
    await page.goto(f"{address}/#/runs/{run.id}")
    layout: Locator = page.locator(".run-layout")
    await expect(page.locator(".answer")).to_contain_text("stay where they are")
    panel: str | None = await layout.get_attribute("data-inspector")
    chips: Locator = page.locator("button.attach-chip")
    await expect(chips).to_have_count(4)
    view: Locator = page.locator(".file-view")

    # A PDF: its pages in their proportions, the first drawn by the daemon.
    report: Locator = chips.filter(has_text="report.pdf")
    await report.click()
    await expect(view).to_have_attribute("data-kind", "pdf")
    await expect(view.locator(".file-page")).to_have_count(2)
    first: Locator = view.locator(".file-page").first
    await expect(first).to_have_attribute("data-drawn", "true")
    await expect(first.locator("img")).to_have_attribute("src", re.compile(r"."))
    assert await first.locator("img").evaluate("(image) => image.decode().then(() => image.naturalWidth)") >= 600
    shown: dict[str, float] | None = await view.bounding_box()
    assert shown is not None
    if width < 960:
        # On a phone the file covers the page.
        assert shown["width"] == width
    await page.screenshot(path=str(shots / f"pdf-{width}-{scheme}.png"))
    await _close(page, view, report)

    # A workbook as Quick Look draws it (item 268): a sheet framed as it is, in a sandboxed frame; Quick Look leaves the
    # hidden sheet out, so the one sheet has no tabs over it. Then what Dot reads of it: a tab for each sheet, the
    # hidden one named so, and what a cell holds.
    budget: Locator = chips.filter(has_text="budget.xlsx")
    await budget.click()
    await expect(view).to_have_attribute("data-kind", "workbook")
    drawn: FrameLocator = view.frame_locator("iframe.file-look")
    await expect(drawn.locator("td")).to_contain_text(["Item", "Amount", "Rent", "1200"])
    await expect(view.locator(".file-tabs[aria-label='Sheets']")).to_have_count(0)
    await expect(view.locator("iframe.file-look")).to_have_attribute("sandbox", "allow-scripts")
    await page.screenshot(path=str(shots / f"workbook-drawn-{width}-{scheme}.png"))
    await view.get_by_role("button", name="Contents").click()
    grid: Locator = view.locator(".file-grid")
    await expect(grid.locator("td")).to_have_text(["Item", "Amount", "Rent", "1200", "Food", "300", "", "=SUM(B2:B3)"])
    await grid.locator("td.pending").click()
    await expect(view.locator(".file-cell-line")).to_have_text("B4=SUM(B2:B3) → not worked out yet")
    await page.screenshot(path=str(shots / f"workbook-{width}-{scheme}.png"))
    await view.get_by_role("button", name="Notes (hidden)").click()
    await expect(grid.locator("td")).to_have_text(["Paid on the first"])
    await _close(page, view, budget)

    # A document as Quick Look draws it, then what Dot reads of it: its heading, its list and its table.
    minutes: Locator = chips.filter(has_text="minutes.docx")
    await minutes.click()
    await expect(view).to_have_attribute("data-kind", "document")
    await expect(view.frame_locator("iframe.file-look").get_by_text("Budget 2027")).to_be_visible()
    await page.screenshot(path=str(shots / f"document-drawn-{width}-{scheme}.png"))
    await view.get_by_role("button", name="Contents").click()
    await expect(view.get_by_role("heading", name="Budget 2027")).to_be_visible()
    await expect(view.locator("li")).to_have_text(["Rent", "Food"])
    await expect(view.locator("tbody td")).to_have_text(["Rent", "1200"])
    await page.screenshot(path=str(shots / f"document-{width}-{scheme}.png"))
    await _close(page, view, minutes)

    # A presentation as Quick Look draws it, then what Dot reads of it: its slide's title, its text and its chart's
    # series by quarter.
    deck: Locator = chips.filter(has_text="deck.pptx")
    await deck.click()
    await expect(view).to_have_attribute("data-kind", "presentation")
    await expect(view.frame_locator("iframe.file-look").get_by_text("Up in the spring")).to_be_visible()
    await page.screenshot(path=str(shots / f"presentation-drawn-{width}-{scheme}.png"))
    await view.get_by_role("button", name="Contents").click()
    slide: Locator = view.locator(".file-slide")
    await expect(slide.get_by_role("heading", name="Sales")).to_be_visible()
    await expect(slide.locator(".file-line")).to_have_text(["Up in the spring"])
    await expect(slide.locator(".file-chart tbody td")).to_have_text(["2027", "1", "2.5"])
    await page.screenshot(path=str(shots / f"presentation-{width}-{scheme}.png"))
    await _close(page, view, deck)

    # A path the answer names: a CSV table as a grid, split at its semicolons, which opens on the Mac.
    paths: Locator = page.locator("button.code-path")
    await expect(paths).to_have_count(5)
    rows: Locator = paths.filter(has_text="rows.csv")
    await rows.click()
    await expect(view).to_have_attribute("data-kind", "table")
    await expect(view.locator(".file-grid td")).to_have_text(["Item", "Amount", "Rent", "1200,50", "Food", "300"])
    await expect(view.get_by_role("button", name="Copy the text")).to_be_hidden()
    await view.get_by_role("button", name="Open on the Mac").click()
    await expect(page.locator(".toast.ok")).to_have_text("Opened rows.csv on the Mac")
    assert opened[-1] == (user_home / "Documents" / "rows.csv").resolve()
    await page.screenshot(path=str(shots / f"table-{width}-{scheme}.png"))
    await _close(page, view, rows)

    # A Markdown file, formatted (item 269), its HTML never drawn, and as it is written a click away.
    todo: Locator = paths.filter(has_text="todo.md")
    await todo.click()
    await expect(view).to_have_attribute("data-kind", "text")
    await expect(view.locator(".file-markdown").get_by_role("heading", name="To do")).to_be_visible()
    await expect(view.locator(".file-markdown p")).to_have_text("Pay the <b>rent</b> first.")
    await expect(view.locator("b")).to_have_count(0)
    await view.get_by_role("button", name="Source").click()
    await expect(view.locator("pre.file-text")).to_have_text(TODO)
    # The text of the file is a press away (item 403); a picture has none.
    await expect(view.get_by_role("button", name="Copy the text")).to_be_visible()
    await _close(page, view, todo)

    # An image, as it is: in a dialog it is a frame around the picture rather than a tall sheet with the picture at its
    # top, and a press shows it at full size over the page.
    chart: Locator = paths.filter(has_text="chart.png")
    await chart.click()
    await expect(view).to_have_attribute("data-kind", "image")
    picture: Locator = view.locator("img.file-image")
    await expect(picture).to_have_attribute("src", re.compile(r"^blob:"))
    assert await picture.evaluate("(image) => image.decode().then(() => image.naturalWidth)") == 320
    await expect(view.get_by_role("button", name="Copy the text")).to_be_hidden()
    if width < 960 or panel != "on":
        frame: dict[str, float] | None = await view.locator(".file-body").bounding_box()
        drawn: dict[str, float] | None = await view.locator(".file-picture").bounding_box()
        assert frame is not None and drawn is not None
        assert frame["height"] - drawn["height"] <= 32, (frame, drawn)
    await view.get_by_role("button", name="chart.png, show at full size").click()
    lightbox: Locator = page.locator("dialog.lightbox")
    await expect(lightbox).to_be_visible()
    await expect(lightbox.locator("img")).to_have_attribute("src", re.compile(r"^blob:"))
    await page.keyboard.press("Escape")
    await expect(lightbox).not_to_be_visible()
    await _close(page, view, chart)

    # A source file as code (item 624): its language and path over it, its lines numbered, its words coloured by kind,
    # and the text whole a press away.
    script: Locator = paths.filter(has_text="total.py")
    await script.click()
    await expect(view).to_have_attribute("data-kind", "text")
    editor: Locator = view.locator("pre.file-code-editor")
    await expect(editor).to_be_visible()
    await expect(view.locator(".file-code-language")).to_have_text("python")
    await expect(view.locator(".file-code-breadcrumb")).to_contain_text("total.py")
    await expect(editor.locator(".file-code-gutter")).to_have_text(["1", "2", "3", "4"])
    await expect(editor.locator(".file-code-keyword").first).to_have_text("def")
    # Python's "//" is division, not a comment (the review of item 624).
    await expect(editor.locator(".file-code-comment")).to_have_count(1)
    await expect(editor.locator(".file-code-comment")).to_have_text("# The sum of the amounts.")
    await expect(editor.locator(".file-code-string")).to_have_text('"amount"')
    await expect(view.get_by_role("button", name="Copy the text")).to_be_visible()
    await _close(page, view, script)

    # A path the panel does not show says why; Escape closes it and the panel is back as it was.
    await paths.filter(has_text="notes.txt").click()
    await expect(view.locator(".note")).to_have_text("This file cannot be shown: the file is in a hidden folder.")
    await page.keyboard.press("Escape")
    await expect(view).to_have_count(0)
    assert await layout.get_attribute("data-inspector") == panel

    # A long text pasted into the box goes as a file, which opens beside the conversation from its chip (item 269).
    box: Locator = page.get_by_role("textbox", name=re.compile(r"^Message for "))
    assert await box.evaluate(
        "(box, text) => { const data = new DataTransfer(); data.setData('text/plain', text);"
        " return !box.dispatchEvent(new ClipboardEvent('paste', { clipboardData: data, bubbles: true, cancelable: "
        "true })); }",
        PASTED,
    )
    pasted: Locator = page.locator(".composer .attach-chip")
    await expect(pasted).to_have_class(re.compile(r"\bready\b"))
    await pasted.locator("button.attach-open").click()
    await expect(view).to_have_attribute("data-kind", "text")
    await expect(view.locator(".file-markdown").get_by_role("heading", name="Pasted plan")).to_be_visible()
    await expect(view.locator(".file-markdown li")).to_have_count(40)
    await expect(page.locator("dialog.file-dialog")).to_have_count(0)
    assert await box.input_value() == ""
    await page.screenshot(path=str(shots / f"pasted-{width}-{scheme}.png"))
    await page.keyboard.press("Escape")
    await expect(view).to_have_count(0)

    # The side panel's Run tab lists no files (item 544): they open from their chips and paths in the conversation.
    if width < 960 or panel != "on":
        await page.get_by_role("button", name="Run details", exact=True).click()
    details: Locator = page.locator(".inspector")
    await details.get_by_role("tab", name="Run").click()
    await expect(details.get_by_role("tabpanel").locator(".panel-section-title").first).to_be_visible()
    await expect(details.locator(".panel-section-title", has_text="Files")).to_have_count(0)

    # The file stays beside the conversation as the chat goes on at its next run (item 406): a message sent from the
    # box opens the next run's page with the file still there, until it is closed. Before, the page made again dropped
    # it.
    if width >= 960:
        await budget.click()
        await expect(view).to_have_attribute("data-kind", "workbook")
        before: str = page.url
        await box.fill("Go on")
        await box.press("Enter")
        await expect(page).not_to_have_url(before)
        await expect(page.locator(".answer").last).to_contain_text("stay where they are")
        await expect(view).to_have_attribute("data-kind", "workbook")
        await view.get_by_role("button", name="Close the file").click()
        await expect(view).to_have_count(0)
        await page.reload()
        await expect(page.locator(".answer").last).to_contain_text("stay where they are")
        await expect(view).to_have_count(0)

    await page.close()


async def _refusals(browser: Browser, address: str, report: StoredAttachment, opened: list[Path]) -> None:
    """The routes find every file with the policy of the panel, and draw a PDF's page as a PNG."""
    page: Page = await browser.new_page()
    headers: dict[str, str] = {"Authorization": f"Bearer {TOKEN}"}
    outside: APIResponse = await page.request.get(f"{address}/api/files/preview?path=/etc/hosts", headers=headers)
    assert outside.status == 422
    assert (await outside.json())["message"] == "the file is outside your home folder"
    script: APIResponse = await page.request.post(
        f"{address}/api/files/open", headers=headers, data={"path": "~/Documents/tidy.sh"}
    )
    assert script.status == 422
    assert (await script.json())["message"] == "a .sh file is not opened in an application"
    assert [path.name for path in opened] == ["rows.csv"] * len(_SHOTS)
    drawn: APIResponse = await page.request.get(
        f"{address}/api/files/page", headers=headers, params={"path": report.path, "number": 1, "width": 600}
    )
    assert drawn.status == 200
    assert drawn.headers["content-type"] == "image/png"
    assert drawn.headers["x-content-type-options"] == "nosniff"
    assert (await drawn.body()).startswith(b"\x89PNG\r\n\x1a\n")
    odd: APIResponse = await page.request.get(
        f"{address}/api/files/page", headers=headers, params={"path": report.path, "number": 1, "width": 601}
    )
    assert odd.status == 422
    await page.close()


@pytest.mark.browser
async def test_a_file_a_message_carries_or_an_answer_names_opens_beside_the_conversation(
    login_env: dict[str, str], tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Before, a file's chip and a path in an answer were text: nothing showed the file."""
    opened: list[Path] = []
    monkeypatch.setattr(routes_files, "open_document", opened.append)
    with scratch_home() as home, eval_environment(home, {}):
        user_home: Path = _user_home(home)
        async with _daemon(login_env, user_home) as (runtime, manager, address), async_playwright() as playwright:
            files: tuple[StoredAttachment, ...] = (
                store_attachment(runtime.workspace, "report.pdf", "application/pdf", _pdf()),
                store_attachment(runtime.workspace, "budget.xlsx", OFFICE_MEDIA_TYPES["workbook"], _workbook()),
                store_attachment(runtime.workspace, "minutes.docx", OFFICE_MEDIA_TYPES["document"], _document()),
                store_attachment(runtime.workspace, "deck.pptx", OFFICE_MEDIA_TYPES["presentation"], _presentation()),
            )
            browser: Browser = await playwright.chromium.launch(channel=runtime.config.browser.channel, headless=True)
            try:
                for size in _SHOTS:
                    await _one_size(browser, runtime, manager, address, files, user_home, opened, size, tmp_path)
                await _refusals(browser, address, files[0], opened)
            finally:
                await browser.close()


@pytest.mark.browser
async def test_a_file_quick_look_cannot_draw_shows_what_dot_reads_of_it_whole(
    login_env: dict[str, str], monkeypatch: pytest.MonkeyPatch
) -> None:
    """Before, the panel kept the layout of a drawing after Quick Look failed, and the contents showed cut off."""
    monkeypatch.setattr(quicklook, "QUICK_LOOK", "/usr/bin/false")
    with scratch_home() as home, eval_environment(home, {}):
        async with (
            _daemon(login_env, _user_home(home)) as (runtime, manager, address),
            async_playwright() as playwright,
        ):
            minutes: StoredAttachment = store_attachment(
                runtime.workspace, "minutes.docx", OFFICE_MEDIA_TYPES["document"], _document()
            )
            browser: Browser = await playwright.chromium.launch(channel=runtime.config.browser.channel, headless=True)
            try:
                page: Page = await browser.new_page(viewport={"width": 1280, "height": 900}, reduced_motion="reduce")
                await page.goto(ui_url("127.0.0.1", int(address.rpartition(":")[2]), TOKEN))
                run: RunRecord = await manager.start(
                    "Look at the minutes",
                    "task",
                    SCRIPTED_MODEL,
                    None,
                    "test",
                    None,
                    None,
                    "standard",
                    None,
                    "ask",
                    False,
                    (minutes,),
                )
                async with asyncio.timeout(WAIT_S):
                    assert (await manager.wait(run.id)).status == "done"
                await page.goto(f"{address}/#/runs/{run.id}")
                await page.locator("button.attach-chip").filter(has_text="minutes.docx").click()
                view: Locator = page.locator(".file-view")
                await expect(view.get_by_text("Quick Look could not draw this file")).to_be_visible()
                await expect(view.locator(".file-body")).to_have_attribute("data-view", "contents")
                await expect(view.get_by_role("heading", name="Budget 2027")).to_be_visible()
            finally:
                await browser.close()
